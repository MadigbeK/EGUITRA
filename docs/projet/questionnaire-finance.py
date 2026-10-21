"""Questionnaire de paramétrage destiné à la direction financière et au comptable.

    python docs/projet/questionnaire-finance.py     # Questionnaire_parametrage_finance.xlsx et .md
    node docs/projet/questionnaire-finance.js       # version Word, depuis le .md

Version courte : la comptabilité est tenue aujourd'hui dans un classeur Excel (le prototype de la
direction) ; les questions partent de ce classeur et des décisions relevées à l'essai de reprise.
Les thèmes sont rangés par urgence : les premiers conditionnent la reprise des données 2026.
Source unique des questions : la liste THEMES ci-dessous.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = Path(__file__).parent
XLSX = HERE / "Questionnaire_parametrage_finance.xlsx"
MD = HERE / "questionnaire-finance.md"

DAF, CPT, DG = "DAF", "Comptable", "Direction"

# (thème, [(question, destinataire, pièce à fournir)])
THEMES = [
    ("Le classeur actuel", [
        ("Quelle est la version définitive du classeur 2026, à quelle date est-elle arrêtée, et qui la tient à jour ?",
         DAF, "Classeur définitif"),
        ("MB AxisPro a-t-elle son propre classeur, ou ses opérations sont-elles dans celui d'EGUITRA GROUP ?",
         DAF, ""),
        ("Tenez-vous d'autres fichiers à côté du classeur : caisse, parc, factures faites sous Word, suivis de chantiers ou de loyers ?",
         CPT, "Ces fichiers"),
        ("Qui établit les états financiers annuels et la déclaration fiscale (DSF) ? Les comptes 2025 sont-ils arrêtés et déposés ?",
         DAF, "États financiers 2025"),
    ]),
    ("Décisions relevées à l'essai de reprise", [
        ("Amortissements : faut-il garder la méthode du classeur (mois entiers à partir du mois suivant l'achat) ou adopter le calcul au jour près ? À trancher avec l'expert-comptable.",
         DAF, ""),
        ("Emprunts : quel est le taux de chaque emprunt ? Le déblocage de mars 2026 semble compté deux fois dans le solde d'ouverture du compte 162 : pouvez-vous le confirmer ?",
         DAF, "Contrats et échéanciers"),
        ("Les seuils repris du classeur vous conviennent-ils (trésorerie minimale, marge transport, délai de recouvrement, couverture de la dette, dépassement de budget) ?",
         DG, ""),
    ]),
    ("Les deux sociétés", [
        ("Raison sociale, adresse, NIF et RCCM de chaque société ?",
         DAF, "RCCM, attestations NIF"),
        ("Quel est le régime fiscal de chaque société, et les deux sont-elles assujetties à la TVA ? Quand a commencé le premier exercice de MB AxisPro ?",
         DAF, ""),
        ("Pour chaque société, qui saisit, qui contrôle, qui approuve et qui paie ?",
         DAF, ""),
        ("Frais communs (bâtiment, personnel, direction, véhicules) : qui les paie, et l'autre société en rembourse-t-elle une part ? Selon quelle clé ?",
         DG, ""),
        ("Y a-t-il des prêts, des avances ou des prestations facturées entre les deux sociétés ?",
         DAF, "Soldes au 31/12/2025"),
    ]),
    ("Plan de comptes et activités", [
        ("Quels sous-comptes utilisez-vous au-delà du plan SYSCOHADA (fournisseurs locaux et étrangers, clients par type, charges détaillées) ?",
         CPT, "Liste des comptes utilisés"),
        ("Quelles activités suivre en résultat ? Pour EGUITRA : transport pétrolier, minier, conteneurs, personnel, location, taxis, BTP, immobilier, négoce, et le pôle agriculture et élevage s'il est actif. Pour MB AxisPro : projets, formation, conseil, audit, certification.",
         DG, ""),
        ("Quels sont vos centres de coût, et comment répartissez-vous les frais généraux entre activités ?",
         DAF, ""),
    ]),
    ("Fiscalité", [
        ("Certaines ventes sont-elles exonérées de TVA ou à un taux particulier (clients miniers, export) ?",
         DAF, "Attestations d'exonération"),
        ("Quelles TVA sur achats ne récupérez-vous pas (carburant, véhicules de tourisme, réception) ?",
         CPT, ""),
        ("Quelles retenues à la source pratiquez-vous sur vos fournisseurs, et lesquelles vos clients opèrent-ils sur vos factures ? Avec quels taux ?",
         DAF, ""),
        ("Quels autres impôts payez-vous (impôt minimum, acomptes d'IS, patente, taxes sur salaires), et à quelles échéances ?",
         DAF, "Dernières déclarations"),
    ]),
    ("Ventes", [
        ("Quel est le dernier numéro de facture émis en 2026, et voulez-vous garder la même numérotation ?",
         CPT, "Une facture récente"),
        ("Quels délais de paiement accordez-vous à vos clients ?",
         DAF, ""),
        ("Transport pétrolier : facturez-vous le volume chargé ou livré, corrigé ou non à 15 °C ? Quel manquant tolérez-vous, et à qui est-il imputé ?",
         DAF, "Grille tarifaire 2026"),
        ("Comment facturez-vous les autres activités : minerais, conteneurs, location, taxis, chantiers (retenue, avance), loyers (TVA, dépôt), honoraires du cabinet ?",
         DAF, "Un exemple de chaque"),
        ("Avez-vous des clients ou des fournisseurs en devises ? Quel taux de change retenez-vous ?",
         DAF, ""),
    ]),
    ("Achats et approbations", [
        ("Quelles sont vos dépenses courantes, et doivent-elles être rattachées à un camion ou à un chantier ? Le carburant est-il acheté par bons, par cartes ou stocké en cuve ?",
         CPT, ""),
        ("Les seuils d'approbation conviennent-ils : finance pour tout achat, direction au-delà de 50 M GNF d'achat et de 100 M GNF de paiement ? Qui remplace le directeur général en son absence ?",
         DG, ""),
        ("Les frais de route des chauffeurs sont-ils avancés puis justifiés, ou payés au forfait ?",
         CPT, ""),
    ]),
    ("Trésorerie", [
        ("Quels sont vos comptes bancaires, vos caisses et vos comptes de mobile money, pour chaque société ?",
         DAF, "RIB et soldes au 31/12/2025"),
        ("Sous quel format obtenez-vous vos relevés bancaires ? Faites-vous aujourd'hui des rapprochements bancaires ?",
         CPT, "Un relevé par banque"),
    ]),
    ("Immobilisations et financements", [
        ("Quelles durées d'amortissement appliquez-vous par catégorie (tracteurs, citernes, véhicules, engins, bâtiments, informatique) ? Avez-vous des contrats de crédit-bail ?",
         CPT, "Tableau des immobilisations"),
        ("Avez-vous des découverts, des cautions bancaires ou des engagements imposés par vos banques (ratio de couverture de la dette) ?",
         DAF, ""),
    ]),
    ("Clôture et pilotage", [
        ("Quels clients sont douteux ou en litige ? À combien de jours de retard relancez-vous ?",
         DAF, "Liste des impayés"),
        ("Arrêtez-vous les comptes chaque mois ? Si oui, à quelle date et avec quelles écritures (charges à payer, provisions) ?",
         CPT, ""),
        ("Quels indicateurs la direction attend-elle chaque mois ? Un budget 2027 est-il prévu ?",
         DG, "Budget 2026"),
        ("À quelle date souhaitez-vous basculer, et combien de temps tenir Excel en parallèle ?",
         DG, ""),
    ]),
    ("Paie et utilisateurs", [
        ("Avec quel outil la paie est-elle faite ? Pouvez-vous fournir chaque mois l'écriture de paie par société ?",
         CPT, "Journal de paie d'un mois"),
        ("Qui utilisera la plateforme, avec quel rôle et pour quelle société ?",
         DG, "Liste nominative"),
    ]),
]

FOREST, SAGE, LINE = "1F3A33", "EEF3F0", "D9E1DC"
THIN = Side(style="thin", color=LINE)
BOTTOM = Border(bottom=THIN)


def numbered():
    n = 0
    for theme, questions in THEMES:
        items = []
        for question, qui, piece in questions:
            n += 1
            items.append((n, question, qui, piece))
        yield theme, items


def xlsx():
    book = Workbook()
    sheet = book.active
    sheet.title = "Questions"
    sheet.sheet_view.showGridLines = False
    widths = {"A": 5, "B": 78, "C": 12, "D": 26, "E": 58}
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width
    sheet["A1"] = "Questionnaire de paramétrage : finance et comptabilité"
    sheet["A1"].font = Font(bold=True, size=15, color=FOREST)
    sheet["A2"] = "EGUITRA GROUP SARLU et MB AxisPro · plateforme NIMOVA"
    sheet["A2"].font = Font(color="7A8A83")
    sheet["A3"] = "Répondez dans la colonne Réponse ; une phrase suffit. Les premiers thèmes conditionnent la reprise des données 2026."
    sheet["A3"].font = Font(italic=True, color="7A8A83")
    row = 5
    for c, title in enumerate(["N°", "Question", "À qui", "Pièce à fournir", "Réponse"], start=1):
        cell = sheet.cell(row=row, column=c, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=FOREST)
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[row].height = 22
    sheet.freeze_panes = "A6"
    total = 0
    for theme, items in numbered():
        row += 1
        cell = sheet.cell(row=row, column=1, value=theme)
        cell.font = Font(bold=True, color=FOREST, size=12)
        for c in range(1, 6):
            sheet.cell(row=row, column=c).fill = PatternFill("solid", fgColor=SAGE)
        sheet.row_dimensions[row].height = 22
        for n, question, qui, piece in items:
            row += 1
            total += 1
            for c, value in enumerate([n, question, qui, piece, ""], start=1):
                cell = sheet.cell(row=row, column=c, value=value)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = BOTTOM
                if c in (1, 3, 4):
                    cell.font = Font(color="5C6B65")
    pieces = book.create_sheet("Pièces à fournir")
    pieces.sheet_view.showGridLines = False
    for col, width in {"A": 44, "B": 12, "C": 14, "D": 16}.items():
        pieces.column_dimensions[col].width = width
    pieces["A1"] = "Pièces à fournir"
    pieces["A1"].font = Font(bold=True, size=15, color=FOREST)
    for c, title in enumerate(["Pièce", "Question", "Remise le", "Par"], start=1):
        cell = pieces.cell(row=3, column=c, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=FOREST)
    r = 3
    for _theme, items in numbered():
        for n, _q, _qui, piece in items:
            if piece:
                r += 1
                for c, value in enumerate([piece, n, "", ""], start=1):
                    cell = pieces.cell(row=r, column=c, value=value)
                    cell.border = BOTTOM
                pieces.cell(row=r, column=3).number_format = "DD/MM/YYYY"
    for ws in (sheet, pieces):
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "5:5"
    book.save(XLSX)
    return total


def markdown():
    out = ["# Questionnaire de paramétrage", "",
           "Finance et comptabilité d'EGUITRA GROUP SARLU et de MB AxisPro. Les réponses fixent le "
           "réglage de la plateforme NIMOVA et la reprise du classeur 2026 ; les premiers thèmes sont les plus urgents. "
           "Une phrase par réponse suffit, de préférence dans le classeur « Questionnaire_parametrage_finance.xlsx ».", ""]
    pieces = []
    for theme, items in numbered():
        out += [f"## {theme}", ""]
        for n, question, qui, piece in items:
            out += [f"{n}. {question} ({qui})", ""]
            if piece:
                pieces.append(f"{piece} (question {n})")
    out += ["## Pièces à fournir", ""] + [f"- {p}" for p in pieces] + [""]
    MD.write_text("\n".join(out), encoding="utf-8", newline="\n")


if __name__ == "__main__":
    total = xlsx()
    markdown()
    print(f"{total} questions : {XLSX.name}, {MD.name}")
