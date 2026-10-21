"""Mise en page commune des questionnaires de paramétrage (classeur Excel et texte pour la version Word).

Utilisé par questionnaire-finance.py et questionnaire-axispro.py : une feuille « Questions » sobre
(thèmes en bandeau, numéro, question, destinataire, pièce, réponse) et une feuille « Pièces à
fournir » à cocher ; le texte Markdown est converti en Word par le script .js du questionnaire.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FOREST, SAGE, LINE, GREY = "1F3A33", "EEF3F0", "D9E1DC", "7A8A83"
BOTTOM = Border(bottom=Side(style="thin", color=LINE))


def numbered(themes):
    """Thèmes avec des questions numérotées de bout en bout : (thème, [(n, question, qui, pièce)])."""
    n = 0
    for theme, questions in themes:
        items = []
        for question, qui, piece in questions:
            n += 1
            items.append((n, question, qui, piece))
        yield theme, items


def _header(ws, row, titles):
    for c, title in enumerate(titles, start=1):
        cell = ws.cell(row=row, column=c, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=FOREST)
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 22


def write_xlsx(path, themes, title, subtitle, note):
    book = Workbook()
    sheet = book.active
    sheet.title = "Questions"
    sheet.sheet_view.showGridLines = False
    for col, width in {"A": 5, "B": 78, "C": 13, "D": 26, "E": 58}.items():
        sheet.column_dimensions[col].width = width
    for cell, text, font in (("A1", title, Font(bold=True, size=15, color=FOREST)),
                             ("A2", subtitle, Font(color=GREY)),
                             ("A3", note, Font(italic=True, color=GREY))):
        sheet[cell] = text
        sheet[cell].font = font
    _header(sheet, 5, ["N°", "Question", "À qui", "Pièce à fournir", "Réponse"])
    sheet.freeze_panes = "A6"
    row, total = 5, 0
    for theme, items in numbered(themes):
        row += 1
        sheet.cell(row=row, column=1, value=theme).font = Font(bold=True, color=FOREST, size=12)
        for c in range(1, 6):
            sheet.cell(row=row, column=c).fill = PatternFill("solid", fgColor=SAGE)
        sheet.row_dimensions[row].height = 22
        for n, question, qui, piece in items:
            row += 1
            total += 1
            for c, value in enumerate([n, question, qui, piece, ""], start=1):
                cell = sheet.cell(row=row, column=c, value=value)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = BOTTOM
                if c in (1, 3, 4):
                    cell.font = Font(color="5C6B65")

    pieces = book.create_sheet("Pièces à fournir")
    pieces.sheet_view.showGridLines = False
    for col, width in {"A": 46, "B": 12, "C": 14, "D": 16}.items():
        pieces.column_dimensions[col].width = width
    pieces["A1"] = "Pièces à fournir"
    pieces["A1"].font = Font(bold=True, size=15, color=FOREST)
    _header(pieces, 3, ["Pièce", "Question", "Remise le", "Par"])
    r = 3
    for _theme, items in numbered(themes):
        for n, _q, _qui, piece in items:
            if piece:
                r += 1
                for c, value in enumerate([piece, n, "", ""], start=1):
                    pieces.cell(row=r, column=c, value=value).border = BOTTOM
                pieces.cell(row=r, column=3).number_format = "DD/MM/YYYY"
    for ws in (sheet, pieces):
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "5:5"
    book.save(path)
    return total


def write_markdown(path, themes, title, intro):
    out = [f"# {title}", "", intro, ""]
    pieces = []
    for theme, items in numbered(themes):
        out += [f"## {theme}", ""]
        for n, question, qui, piece in items:
            out += [f"{n}. {question} ({qui})", ""]
            if piece:
                pieces.append(f"{piece} (question {n})")
    if pieces:
        out += ["## Pièces à fournir", ""] + [f"- {p}" for p in pieces] + [""]
    path.write_text("\n".join(out), encoding="utf-8", newline="\n")
