"""Génère docs/projet/Fiche_de_retours.xlsx : retours des testeurs de la recette (une ligne par remarque).

    python docs/projet/fiche-de-retours.py

Les testeurs remplissent les colonnes de gauche (compte, application, écran, type, description,
gravité) ; E-VOLUTION XP remplit la réponse et le statut. L'onglet Synthèse se calcule seul.
Accompagne le guide des testeurs (guide-des-testeurs.md).
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FOREST, GREEN, GREY = "142F29", "22634C", "F2F4F2"
OUT = Path(__file__).with_name("Fiche_de_retours.xlsx")
LIGNES = 300

COMPTES = ["test.dg", "test.daf", "test.comptable", "test.exploitation", "test.parc", "test.chantiers",
           "test.immobilier", "test.consultant", "test.admin", "test.porteur"]
APPLICATIONS = ["Finance", "AxisPro", "Rotations", "Portail", "Connexion"]
TYPES = ["Anomalie", "Amélioration", "Manque", "Question"]
GRAVITES = ["Bloquant", "Gênant", "Mineur"]
STATUTS = ["Nouveau", "Pris en compte", "Corrigé en recette", "Mis en production", "Hors périmètre, à arbitrer",
           "Sans suite"]
# (titre, largeur, rempli par)
COLONNES = [
    ("N°", 6, "testeur"), ("Date", 12, "testeur"), ("Compte", 17, "testeur"), ("Application", 13, "testeur"),
    ("Écran (adresse)", 30, "testeur"), ("Type", 14, "testeur"), ("Ce que j'ai fait", 40, "testeur"),
    ("Ce que j'attendais", 34, "testeur"), ("Ce qui s'est passé", 34, "testeur"), ("Gravité", 11, "testeur"),
    ("Capture (nom du fichier)", 20, "testeur"), ("Réponse d'E-VOLUTION XP", 40, "prestataire"),
    ("Statut", 22, "prestataire"), ("Date de traitement", 14, "prestataire"),
]

THIN = Side(style="thin", color="D5DBD7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = Font(bold=True, color="FFFFFF")
FILL_HEAD = PatternFill("solid", fgColor=FOREST)
FILL_PRESTA = PatternFill("solid", fgColor=GREEN)
FILL_ALT = PatternFill("solid", fgColor=GREY)
WRAP = Alignment(wrap_text=True, vertical="top")

book = Workbook()

# ─── Mode d'emploi ────────────────────────────────────────────────────────
guide = book.active
guide.title = "Mode d'emploi"
lignes = [
    ("Fiche de retours des testeurs : plateforme NIMOVA (recette)", Font(bold=True, size=16, color=FOREST)),
    ("EGUITRA GROUP SARLU et MB AxisPro · E-VOLUTION XP", Font(italic=True, color="666666")),
    ("", None),
    ("Comment remplir", Font(bold=True, color=GREEN)),
    ("Une ligne par remarque, dans l'onglet Retours. Choisissez le compte, l'application, le type et la gravité "
     "dans les listes ; copiez l'adresse de l'écran depuis la barre d'adresse du navigateur.", None),
    ("Décrivez ce que vous avez fait, ce que vous attendiez et ce qui s'est passé : c'est ce qui permet de "
     "reproduire et de corriger.", None),
    ("Capture d'écran : Windows + Maj + S, puis enregistrez-la et indiquez son nom dans la colonne prévue.", None),
    ("", None),
    ("Types", Font(bold=True, color=GREEN)),
    ("Anomalie : quelque chose ne marche pas. Amélioration : cela marche mais pourrait mieux faire. Manque : une "
     "fonction attendue n'existe pas. Question : un point à éclaircir.", None),
    ("Gravité : Bloquant (impossible de continuer), Gênant, Mineur.", None),
    ("", None),
    ("Envoi", Font(bold=True, color=GREEN)),
    ("Envoyez la fiche à la direction en fin de semaine. Un point bloquant se signale tout de suite, par "
     "téléphone ou par message.", None),
    ("Les colonnes vertes foncées (réponse, statut) sont remplies par E-VOLUTION XP.", None),
]
for r, (text, font) in enumerate(lignes, start=1):
    cell = guide.cell(row=r, column=1, value=text)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    if font:
        cell.font = font
guide.column_dimensions["A"].width = 110

# ─── Retours ──────────────────────────────────────────────────────────────
sheet = book.create_sheet("Retours")
for c, (titre, largeur, qui) in enumerate(COLONNES, start=1):
    cell = sheet.cell(row=1, column=c, value=titre)
    cell.font, cell.border = HEAD, BORDER
    cell.fill = FILL_PRESTA if qui == "prestataire" else FILL_HEAD
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.column_dimensions[get_column_letter(c)].width = largeur
sheet.row_dimensions[1].height = 32
for r in range(2, LIGNES + 2):
    sheet.cell(row=r, column=1, value=r - 1)
    for c in range(1, len(COLONNES) + 1):
        cell = sheet.cell(row=r, column=c)
        cell.border, cell.alignment = BORDER, WRAP
        if r % 2 == 1:
            cell.fill = FILL_ALT
    sheet.cell(row=r, column=2).number_format = "DD/MM/YYYY"
    sheet.cell(row=r, column=14).number_format = "DD/MM/YYYY"


def liste(valeurs, colonne):
    dv = DataValidation(type="list", formula1='"' + ",".join(valeurs) + '"', allow_blank=True)
    dv.error, dv.errorTitle = "Choisissez une valeur de la liste.", "Valeur inconnue"
    sheet.add_data_validation(dv)
    lettre = get_column_letter(colonne)
    dv.add(f"{lettre}2:{lettre}{LIGNES + 1}")


liste(COMPTES, 3)
liste(APPLICATIONS, 4)
liste(TYPES, 6)
liste(GRAVITES, 10)
liste(STATUTS, 13)
dates = DataValidation(type="date", operator="greaterThan", formula1="DATE(2026,1,1)", allow_blank=True)
sheet.add_data_validation(dates)
dates.add(f"B2:B{LIGNES + 1}")
sheet.freeze_panes = "C2"
sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLONNES))}{LIGNES + 1}"

# ─── Synthèse ─────────────────────────────────────────────────────────────
summary = book.create_sheet("Synthèse")
summary.cell(row=1, column=1, value="Synthèse des retours").font = Font(bold=True, size=14, color=FOREST)
plage = f"$%s$2:$%s${LIGNES + 1}"


def tableau(ligne, titre, valeurs, colonne_source):
    col = get_column_letter(colonne_source)
    entetes = [titre, "Total", "Bloquant", "Gênant", "Mineur", "Traités"]
    for c, t in enumerate(entetes, start=1):
        cell = summary.cell(row=ligne, column=c, value=t)
        cell.font, cell.fill, cell.border = HEAD, FILL_HEAD, BORDER
    source, grav, statut = plage % (col, col), plage % ("J", "J"), plage % ("M", "M")
    for i, v in enumerate(valeurs, start=1):
        r = ligne + i
        cellules = [v, f'=COUNTIF(Retours!{source},A{r})'] + [
            f'=COUNTIFS(Retours!{source},A{r},Retours!{grav},"{g}")' for g in GRAVITES
        ] + [f'=COUNTIFS(Retours!{source},A{r},Retours!{statut},"Corrigé en recette")'
             f'+COUNTIFS(Retours!{source},A{r},Retours!{statut},"Mis en production")']
        for c, value in enumerate(cellules, start=1):
            cell = summary.cell(row=r, column=c, value=value)
            cell.border = BORDER
    return ligne + len(valeurs) + 2


suite = tableau(3, "Type", TYPES, 6)
suite = tableau(suite, "Application", APPLICATIONS, 4)
tableau(suite, "Compte", COMPTES, 3)
summary.column_dimensions["A"].width = 22
for c in range(2, 7):
    summary.column_dimensions[get_column_letter(c)].width = 12

book.save(OUT)
print(f"Écrit : {OUT}")
