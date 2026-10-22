"""Génère docs/projet/Cahier_de_recette.xlsx : scénarios de recette à jouer avec les référents.

    python docs/projet/cahier-de-recette.py

Les scénarios suivent les vingt attentes du département Finance (offre, chapitre 4), AxisPro
Suite (chapitre 3.7) et la reprise (chapitre 6.3). Les référents notent le résultat obtenu, le
statut (OK, KO, Réserve), leurs initiales et la date ; l'onglet Synthèse se calcule seul.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY, TEAL, GREY = "1C1E6B", "1A9E9E", "F2F4F7"
OUT = Path(__file__).with_name("Cahier_de_recette.xlsx")

ATTENTES = {
    "1": "Visibilité temps réel : CA, résultat, trésorerie",
    "2": "Vue consolidée, multi-sociétés, décomposable",
    "3": "Comptabilité SYSCOHADA révisé",
    "4": "Analytique par centre de coût, chantier, activité",
    "5": "Trésorerie et rapprochement bancaire",
    "6": "Immobilisations et amortissements",
    "7": "Budget, écarts réalisé et prévisionnel",
    "8": "Facturation clients et fournisseurs",
    "9": "Fiscalité, TVA, déclarations",
    "10": "Approbation des engagements et paiements",
    "11": "Coûts et rentabilité transport",
    "12": "Chantiers, avancement, retenues",
    "13": "Gestion locative et patrimoine",
    "14": "Traçabilité, piste d'audit",
    "15": "Sécurité des accès par profil",
    "16": "Disponibilité et sauvegardes",
    "17": "Interface simple et formation",
    "18": "Support et accompagnement",
    "19": "Solution évolutive",
    "AX": "AxisPro Suite et portail des porteurs",
}

# (n°, attente, profil, écran, scénario, données d'essai, résultat attendu)
SCENARIOS = [
    ("R01", "1", "Direction", "Finance > Tableau de bord", "Noter le CA du mois ; faire saisir et valider une facture de vente par la finance ; rouvrir le tableau de bord.", "Facture de 10 000 000 GNF HT", "Le CA du mois augmente de 10 000 000 GNF ; l'encours client augmente du montant TTC."),
    ("R02", "1", "Direction", "Finance > Rapport mensuel", "Choisir le mois écoulé, imprimer le rapport.", "—", "Rapport complet : indicateurs, résultat du mois et cumulé comparés à N-1, activités, flux, trésorerie, transport, alertes."),
    ("R03", "2", "Direction (compte ouvert aux deux sociétés)", "En-tête : sélecteur de société", "Passer d'EGUITRA GROUP SARLU à MB AxisPro, puis revenir.", "—", "L'écran se recharge ; seules les pièces de la société choisie apparaissent, jamais celles de l'autre."),
    ("R04", "2", "Direction", "Tableau de bord : « Groupe consolidé »", "Comparer les indicateurs en vue Société et en vue Groupe.", "—", "Le groupe réunit les deux sociétés ; chaque société garde ses propres montants."),
    ("R05", "2", "Administration", "Administration > Utilisateurs", "Créer un utilisateur ouvert aux deux sociétés et un autre ouvert au cabinet seulement ; se connecter avec chacun.", "Profil finance", "Le premier voit le sélecteur et les deux sociétés ; le second ne voit que MB AxisPro."),
    ("R06", "2", "Finance", "Activités et budget", "Afficher la marge par activité, chantier, ensemble routier et route.", "Données reprises", "Montants cohérents avec les écritures imputées."),
    ("R07", "3", "Finance", "Balance et états > Bilan", "Arrêté au dernier jour du mois écoulé.", "—", "Bilan équilibré (actif = passif) ; aucun compte « hors états »."),
    ("R08", "3", "Finance", "Balance et états > Compte de résultat", "Exercice en cours, comparaison à N-1.", "—", "Soldes intermédiaires (marge, valeur ajoutée, EBE, résultat) conformes à la balance."),
    ("R09", "3", "Finance", "Balance et états > Tableau des flux", "Exercice en cours.", "—", "Le tableau rejoint la trésorerie nette du bilan (écart nul)."),
    ("R10", "3", "Finance et expert-comptable", "Balance et états > Balance", "Comparer la balance d'ouverture reprise à la balance de clôture de l'ancien système.", "Balance au 31/12/2025", "Soldes identiques compte par compte, pour chaque société."),
    ("R11", "3", "Finance", "Reprise de l'exercice 2026", "Contrôler les écritures et les rotations reprises.", "Classeur 2026", "Balances mensuelles, nombre de rotations et chiffre d'affaires conformes au classeur."),
    ("R12", "4", "Exploitation", "Transport", "Valider puis facturer une rotation.", "—", "La facture porte l'activité transport, l'ensemble routier et la route ; la marge apparaît par ensemble et par route."),
    ("R13", "4", "Finance", "Achats et dépenses", "Saisir une dépense de carburant imputée à un ensemble routier.", "—", "Le coût apparaît dans la rentabilité de l'ensemble."),
    ("R14", "5", "Finance", "Trésorerie", "Importer un relevé bancaire réel (CSV, XLSX ou camt.053 XML) ; réimporter le même fichier.", "Relevé d'un mois", "Lignes importées, soldes initial et final lus et contrôlés (écart refusé sauf confirmation), doublons ignorés à la réimportation."),
    ("R15", "5", "Finance", "Trésorerie > Rapprochement", "Rapprocher les lignes avec les factures et règlements proposés.", "—", "Lignes rapprochées, factures soldées, écart nul."),
    ("R16", "5", "Finance", "Trésorerie > Virement interne", "Virer entre deux comptes bancaires.", "5 000 000 GNF", "Passage par le compte 585, soldé après les deux mouvements."),
    ("R17", "6", "Finance", "Immobilisations", "Créer une catégorie et une immobilisation, la mettre en service.", "Camion, 5 ans", "Plan d'amortissement mensuel calculé."),
    ("R18", "6", "Finance", "Immobilisations", "Le lendemain, consulter le plan.", "—", "Dotations échues comptabilisées automatiquement."),
    ("R19", "6", "Finance", "Immobilisations > Sortir de l'actif", "Céder une immobilisation.", "Prix de cession", "Écritures de sortie (28/2x, 81x) et de cession (485/82x) ; immobilisation sortie du registre."),
    ("R20", "7", "Finance", "Activités et budget > Budget", "Saisir un budget mensuel par compte ; comparer au réalisé.", "—", "Écarts par compte et par mois ; alerte au-delà du seuil de dépassement."),
    ("R21", "8", "Finance", "Ventes", "Créer, valider et imprimer une facture ; créer un avoir partiel.", "—", "Numérotation propre à la société ; PDF correct ; reste dû diminué de l'avoir."),
    ("R22", "8", "Finance", "Achats et dépenses", "Saisir une facture fournisseur avec sa pièce jointe numérisée ; la soumettre.", "Montant supérieur au seuil", "La pièce part au circuit d'approbation (voir R26)."),
    ("R23", "8", "Finance", "Relances clients", "Calculer les relances ; ouvrir une mise en demeure ; l'imprimer.", "Facture échue depuis plus de 30 jours", "Niveau correct, lettre complète ; la mise en demeure n'est jamais envoyée sans décision."),
    ("R24", "9", "Finance et expert-comptable", "Fiscalité", "Afficher la TVA du mois.", "—", "TVA collectée et déductible égales aux comptes 443 et 445 ; modèle de déclaration validé."),
    ("R25", "9", "Finance", "Fiscalité > Export d'audit", "Exporter les écritures de l'exercice.", "—", "Fichier complet, ouvert sans erreur par l'expert-comptable."),
    ("R26", "10", "Finance puis direction", "À valider", "Soumettre une pièce d'achat au-delà de 50 M GNF ; approuver par la finance, puis par la direction.", "—", "Pièce comptabilisée seulement après la dernière approbation ; décisions journalisées."),
    ("R27", "10", "Direction", "À valider", "Refuser une demande.", "—", "Motif obligatoire ; la pièce revient au demandeur sans être comptabilisée."),
    ("R28", "10", "Chantiers", "À valider", "Consulter la liste des demandes.", "—", "Seules les demandes qui concernent ce profil apparaissent."),
    ("R29", "11", "Exploitation", "Application Rotations (téléphone)", "Installer l'application sur l'écran d'accueil ; se connecter.", "Téléphone Android", "Icône Rotations ; connexion avec double authentification."),
    ("R30", "11", "Exploitation", "Application Rotations, mode avion", "Saisir deux rotations sans réseau, puis réactiver le réseau.", "—", "Rotations gardées « À envoyer », transmises au retour du réseau, sans doublon."),
    ("R31", "11", "Exploitation", "Transport", "Valider les rotations reçues ; facturer un client.", "—", "Une facture par client, une ligne par rotation ; tarif du client appliqué."),
    ("R32", "11", "Direction", "Transport > Rentabilité", "Marge par ensemble routier, par route et par mois.", "—", "Marges cohérentes ; alerte si la marge passe sous le seuil."),
    ("R33", "12", "Chantiers", "Chantiers", "Créer un chantier ; saisir une situation en pourcentage ; la facturer.", "Marché, retenue de 5 %", "Facture de la période ; retenue portée en 4117."),
    ("R34", "12", "Chantiers et finance", "Chantiers", "Libérer les retenues à la réception.", "—", "Retenues redevenues exigibles du client."),
    ("R35", "13", "Immobilier", "Immobilier", "Créer un bien et un bail ; activer le bail ; revenir le lendemain.", "Loyer mensuel, prise d'effet passée", "Chaque nuit, loyers échus depuis la prise d'effet facturés et validés sans intervention (une facture par échéance, datée de l'échéance)."),
    ("R36", "13", "Immobilier", "Immobilier > Impayés et quittances", "Encaisser un loyer ; imprimer la quittance ; consulter les impayés.", "—", "Quittance délivrée seulement pour un loyer réglé ; impayés à jour."),
    ("R37", "14", "Administration", "Journal d'audit", "Modifier une fiche client, puis consulter le journal.", "—", "Utilisateur, date, valeurs avant et après."),
    ("R38", "14", "Finance", "Clôtures et contrôles", "Clôturer un mois ; tenter une saisie dans le mois clôturé ; rouvrir.", "—", "Contrôles exécutés ; saisie refusée dans la période close ; réouverture tracée au registre."),
    ("R39", "15", "Tout utilisateur", "Connexion", "Première connexion avec le mot de passe provisoire.", "—", "Activation de la double authentification puis choix du mot de passe, obligatoires avant tout accès."),
    ("R40", "15", "Tout utilisateur", "Connexion", "Saisir cinq codes faux de suite.", "—", "Accès bloqué temporairement ; tentatives journalisées."),
    ("R41", "15", "Exploitation", "Navigation", "Tenter d'ouvrir le tableau de bord du dirigeant.", "—", "Accès refusé : le profil ne le permet pas."),
    ("R42", "15", "Administration", "Administration > Utilisateurs", "Désactiver un compte.", "—", "Sessions fermées immédiatement ; historique conservé."),
    ("R43", "16", "Prestataire, devant le client", "Serveur", "Restaurer la production à un instant choisi dans un environnement jetable (backup/restore-test.sh --target … --report).", "Heure choisie par le client", "Base restaurée et contrôlée (sociétés, pièces) ; compte rendu remis ; disponibilité du mois écoulé calculée (availability-report.py)."),
    ("R44", "17", "Chaque profil", "Guide du profil", "Réaliser les tâches du guide sans assistance.", "Guide remis en formation", "Tâches réalisées ; remarques notées."),
    ("R45", "18", "Référent", "Suivi des demandes", "Déclarer une anomalie de recette selon la procédure de support.", "Procédure de support remise", "Accusé de réception, gravité confirmée, prise en compte dans les délais de l'offre (chapitre 7.2)."),
    ("R46", "19", "Prestataire", "Recette puis production", "Ajouter une fonction ou un module du noyau selon le dossier technique, l'installer en recette, la faire valider, la mettre en production.", "Dossier technique", "Procédure de mise en production suivie ; interruption limitée à quelques minutes, annoncée."),
    ("R47", "10", "Finance", "Achats et dépenses", "Remettre en brouillon une pièce approuvée ; changer le fournisseur ; la valider.", "Pièce approuvée", "Validation refusée : la pièce modifiée repart au circuit d'approbation."),
    ("R48", "15", "Exploitation", "Application Rotations", "Se déconnecter ; passer en mode avion ; rouvrir l'application.", "Saisie en attente", "L'application ne s'ouvre plus sous l'identité de l'agent parti ; sa saisie en attente reste et part à sa prochaine connexion."),
    ("R49", "13", "Immobilier", "Immobilier > Biens et baux", "Mettre fin à un bail ; revenir le lendemain ; modifier la valeur actuelle d'un bien.", "Bail en cours", "Loyers dus jusqu'à la fin émis, plus aucun ensuite ; bien libre ; valeur et rendement du patrimoine mis à jour."),
    ("R50", "14", "Finance", "Clôtures et contrôles", "Après une clôture mensuelle, consulter les contrôles d'intégrité.", "Mois clôturé", "Contrôles du mois écoulé et chaînage « Intactes » pour chaque journal verrouillé."),
    ("R51", "2", "Direction", "Balance et états > Consolidé du groupe", "Facturer une prestation d'une société du groupe à l'autre ; afficher le consolidé de l'exercice.", "Facture intra-groupe", "Compte de résultat et bilan par société ; la vente et l'achat réciproques, la créance et la dette s'éliminent ; colonne Groupe = somme des sociétés − éliminations."),
    ("R52", "1", "Direction", "Finance > Tableau de bord", "Choisir une date d'arrêté ; lire les onze questions du dirigeant et le bloc « Ce qui demande une décision » ; imprimer.", "Données reprises", "Chaque question a sa réponse, son verdict (vert, orange, rouge) et sa lecture ; décisions reliées aux alertes."),
    ("R53", "1", "Finance", "Pilotage > Alertes", "Afficher toutes les règles, conformes comprises ; modifier un seuil (DSCR, délai de recouvrement).", "—", "Gravité, valeur, seuil et action recommandée pour chaque règle ; seuil laissé à zéro signalé en vigilance."),
    ("R54", "4", "Finance", "Achats et dépenses", "Saisir une dépense par sa catégorie, sur deux axes (activité et ensemble routier).", "Catégorie Carburant", "Compte de charge, compte du fournisseur et TVA déduits de la catégorie ; ventilation analytique sur les deux axes."),
    ("R55", "7", "Finance", "Activités et budget > Budget", "Importer le budget par ligne budgétaire (XLSX) ; comparer au réalisé et aux engagements.", "Budget 2026 du classeur", "Budget cumulé, réalisé, engagé, disponible, écart et statut (dépassement, conforme, sous-réalisation) par ligne."),
    ("R56", "10", "Chantiers puis finance et direction", "Engagements", "Créer un engagement de dépense de 60 M GNF avec son devis ; le faire approuver ; saisir ensuite la facture du fournisseur.", "Devis joint", "Engagement approuvé par la finance puis la direction ; la facture consomme l'engagement sans refaire le circuit ; reste à engager à jour."),
    ("R57", "10", "Finance puis direction", "Trésorerie", "Saisir une sortie de banque de 150 M GNF ; tenter de la comptabiliser.", "—", "Opération en brouillon jusqu'à l'accord de la direction, comptabilisée à la dernière approbation."),
    ("R58", "3", "Finance", "Clôtures et contrôles > Clôture mensuelle", "Clôturer un mois complet.", "Mois repris", "Dotations, liquidation de la TVA (pièce TVA du mois), factures à établir des rotations non facturées (extournées le mois suivant), sceau et procès-verbal imprimé."),
    ("R59", "3", "Finance et responsable financier", "Clôtures et contrôles > Validation de la reprise", "Charger les valeurs de référence du classeur ; lancer le contrôle ; imprimer l'annexe.", "Valeurs de référence du classeur", "26 indicateurs recalculés, écart de chacun ; « 26/26 » conformes attendu avant la bascule ; annexe jointe au procès-verbal de bascule."),
    ("R60", "11", "Direction", "Transport > Rentabilité", "Afficher la rentabilité par ensemble routier avec les charges de structure ; contrôler les manquants d'une rotation.", "Rotations reprises", "Coût et CA au km, structure, résultat net par ensemble ; manquant toléré, excédentaire déduit ou imputé au chauffeur ; taxes spécifiques au barème."),
    ("R61", "9", "Finance", "Fiscalité", "Saisir une facture fournisseur soumise à retenue à la source ; consulter la déclaration du mois.", "Prestataire non résident ou retenue BIC", "Retenue calculée et portée au compte de l'État ; déclaration et historique à jour."),
    ("R63", "11", "Finance", "Transport > Routes et tarifs", "Renseigner les forfaits d'une route et le prix du carburant ; saisir une rotation, la valider, puis changer le prix.", "Route et ensemble existants", "Coût standard et marge estimée affichés à la saisie ; après validation, un nouveau prix ne change pas la rotation ; rentabilité : coûts standard, déclarés et comptabilisés côte à côte."),
    ("R64", "11", "Finance", "Transport > Flotte, Propriétaires", "Déclarer un ensemble géré pour un propriétaire avec ses charges mensuelles ; créer une règle de 5 % pour la société et une commission de partenaire ; établir les relevés du mois.", "Rotations validées sur l'ensemble", "Relevé : chiffre d'affaires, charges, marge, commissions, part du propriétaire ; PDF et Excel ; ensemble à l'arrêt : part négative."),
    ("R65", "11", "Finance", "Transport > Propriétaires", "Préparer les pièces d'un relevé, les consulter dans Achats et Ventes, puis rouvrir le relevé.", "Relevé établi", "Facture d'achat du propriétaire (ou refacturation d'une part négative) et facture du partenaire en brouillon, imputées sur l'ensemble ; aucune pièce pour la part de la société ; rouvrir supprime les brouillons."),
    ("R62", "5", "Finance", "Pilotage > Trésorerie", "Afficher les flux mensuels et la prévision ; régler une échéance d'emprunt depuis une ligne de relevé.", "Emprunt repris", "Encaissements, décaissements, cash net et solde par mois face au seuil ; échéance proposée au rapprochement, capital et intérêts comptabilisés une seule fois."),
    ("A01", "AX", "Cabinet", "AxisPro > Grille d'évaluation", "Importer la grille officielle ; la vérifier ; la mettre en vigueur.", "Classeur du cabinet", "8 gates et leurs règles, 15 domaines, 111 critères (9 knock-outs), 42 pièces et 102 rattachements chargés ; version en vigueur."),
    ("A02", "AX", "Cabinet", "AxisPro > Dossiers", "Ouvrir un dossier.", "Projet réel anonymisé", "Référence attribuée ; critères du premier gate à évaluer."),
    ("A03", "AX", "Cabinet", "Dossier > Évaluation", "Évaluer les critères en saisie guidée ; déclarer disponible un knock-out sans pièce, puis lui joindre une preuve avec sa date de validité.", "—", "Sans pièce, critère retenu « incomplet » (saisie inchangée) ; avec la preuve, retenu disponible ; Readiness, APD et Bankability recalculés ; pièce de la data room disponible."),
    ("A04", "AX", "Cabinet", "Dossier > Synthèse", "Contrôler les knock-outs et les actions prioritaires ; passer au gate suivant ; forcer un passage refusé avec un motif.", "—", "NO-GO si un knock-out est actif ; passage refusé avec ses motifs ; passage forcé inscrit au journal."),
    ("A05", "AX", "Cabinet", "Dossier > Comités", "Sceller et verrouiller une décision HOLD motivée ; vérifier l'empreinte ; tenter de modifier le dossier ; ouvrir une nouvelle révision ; comparer deux revues.", "—", "Décision non modifiable, empreinte valide ; dossier figé jusqu'à la nouvelle révision ; comparaison des scores, gates et domaines."),
    ("A06", "AX", "Cabinet", "Dossier > Rapport", "Télécharger le rapport du comité (PDF), le classeur d'évaluation, la fiche d'évaluation (la modifier puis la réimporter), la data room (.zip) et les pièces à fournir.", "—", "Politique appliquée et empreinte rappelées ; classeur au plan du modèle du cabinet ; import précédé d'un aperçu des changements."),
    ("A07", "AX", "Cabinet", "Dossier > Temps et facturation", "Saisir des temps ; facturer les honoraires.", "Taux horaire", "Facture d'honoraires en brouillon dans la comptabilité du cabinet."),
    ("A08", "AX", "Porteur de projet", "Portail", "Première connexion ; déposer une pièce ; écrire au cabinet.", "Accès ouvert par le cabinet", "Pièce visible du cabinet ; notes internes du cabinet invisibles du porteur."),
    ("A09", "AX", "Cabinet", "Dossier > Data room", "Valider une pièce ; en refuser une autre avec un motif.", "—", "Le porteur voit l'état de chaque pièce et le motif du refus."),
    ("A10", "AX", "Porteur de projet", "Portail > Pièces", "Déposer un programme renommé en .pdf, puis un document Word à macros.", "Fichiers d'essai fournis par le prestataire", "Dépôts refusés, formats acceptés rappelés (PDF, JPEG, PNG, DOCX, XLSX)."),
    ("A11", "AX", "Direction (société du groupe seule)", "AxisPro Suite", "Ouvrir AxisPro Suite avec un compte de direction ouvert à la seule société du groupe.", "—", "Aucun dossier, aucune pièce, revue ni grille du cabinet visible."),
]

COLUMNS = [("N°", 7), ("Attente", 9), ("Profil", 18), ("Écran", 24), ("Scénario", 46), ("Données d'essai", 20),
           ("Résultat attendu", 46), ("Résultat obtenu", 32), ("Statut", 10), ("Testé par", 12), ("Date", 11), ("Commentaire", 30)]

thin = Side(style="thin", color="C9CFD8")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = Font(bold=True, color="FFFFFF")
FILL_HEAD = PatternFill("solid", fgColor=NAVY)
FILL_ALT = PatternFill("solid", fgColor=GREY)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()

# ─── Mode d'emploi ───────────────────────────────────────────────────────
guide = wb.active
guide.title = "Mode d'emploi"
lines = [
    ("Cahier de recette : EGUITRA Finance et AxisPro Suite", Font(bold=True, size=16, color=NAVY)),
    ("Offre OTF-EGUITRA-2026-37 · E-VOLUTION XP", Font(italic=True, color="666666")),
    ("", None),
    ("Objet", Font(bold=True, color=TEAL)),
    ("Vérifier avec les référents, sur l'environnement de recette, chaque attente de l'offre avant la mise en production.", None),
    ("", None),
    ("Où et avec quels comptes", Font(bold=True, color=TEAL)),
    ("Environnement de recette (adresse remise par le prestataire), avec les comptes personnels de recette de chaque référent.", None),
    ("Les données de recette sont une copie ou un extrait des données réelles : rien de ce qui est saisi en recette n'atteint la production.", None),
    ("", None),
    ("Comment remplir l'onglet Scénarios", Font(bold=True, color=TEAL)),
    ("1. Jouer le scénario tel qu'il est décrit, avec les données d'essai indiquées.", None),
    ("2. Noter le résultat obtenu, puis choisir le statut : OK, KO ou Réserve (écart mineur accepté).", None),
    ("3. Indiquer ses initiales et la date. En cas de KO, décrire l'écart dans le commentaire : le prestataire corrige, le scénario est rejoué.", None),
    ("", None),
    ("Décision de mise en production", Font(bold=True, color=TEAL)),
    ("La mise en production se décide quand tous les scénarios sont OK ou en réserve acceptée par le comité de pilotage ; l'onglet Synthèse fait le compte.", None),
    ("Le procès-verbal de recette reprend la synthèse et la liste des réserves.", None),
]
for row, (text, font) in enumerate(lines, start=1):
    cell = guide.cell(row=row, column=1, value=text)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    if font:
        cell.font = font
guide.column_dimensions["A"].width = 120

# ─── Scénarios ───────────────────────────────────────────────────────────
sheet = wb.create_sheet("Scénarios")
for col, (name, width) in enumerate(COLUMNS, start=1):
    cell = sheet.cell(row=1, column=col, value=name)
    cell.font, cell.fill, cell.border = HEAD, FILL_HEAD, BORDER
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.column_dimensions[get_column_letter(col)].width = width
for r, scenario in enumerate(SCENARIOS, start=2):
    for col, value in enumerate(list(scenario) + [""] * 5, start=1):
        cell = sheet.cell(row=r, column=col, value=value)
        cell.alignment, cell.border = WRAP, BORDER
        if r % 2:
            cell.fill = FILL_ALT
last = len(SCENARIOS) + 1
statuses = DataValidation(type="list", formula1='"OK,KO,Réserve"', allow_blank=True)
statuses.error = "Choisir OK, KO ou Réserve."
sheet.add_data_validation(statuses)
statuses.add(f"I2:I{last}")
sheet.freeze_panes = "E2"
sheet.auto_filter.ref = f"A1:L{last}"
sheet.page_setup.orientation = "landscape"
sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
sheet.sheet_properties.pageSetUpPr.fitToPage = True
sheet.print_title_rows = "1:1"

# ─── Synthèse ────────────────────────────────────────────────────────────
summary = wb.create_sheet("Synthèse")
heads = ["Attente", "Libellé", "Scénarios", "OK", "Réserve", "KO", "Non joués"]
widths = [9, 52, 11, 8, 9, 8, 11]
for col, (name, width) in enumerate(zip(heads, widths, strict=True), start=1):
    cell = summary.cell(row=1, column=col, value=name)
    cell.font, cell.fill, cell.border = HEAD, FILL_HEAD, BORDER
    summary.column_dimensions[get_column_letter(col)].width = width
rng = f"'Scénarios'!$B$2:$B${last}"
st = f"'Scénarios'!$I$2:$I${last}"
for r, (code, label) in enumerate(ATTENTES.items(), start=2):
    key = f'"{code}"'
    values = [code, label, f"=COUNTIF({rng},{key})", f'=COUNTIFS({rng},{key},{st},"OK")',
              f'=COUNTIFS({rng},{key},{st},"Réserve")', f'=COUNTIFS({rng},{key},{st},"KO")', f"=C{r}-D{r}-E{r}-F{r}"]
    for col, value in enumerate(values, start=1):
        cell = summary.cell(row=r, column=col, value=value)
        cell.border = BORDER
        if r % 2:
            cell.fill = FILL_ALT
total = len(ATTENTES) + 2
summary.cell(row=total, column=2, value="Total").font = Font(bold=True)
for col in range(3, 8):
    letter = get_column_letter(col)
    cell = summary.cell(row=total, column=col, value=f"=SUM({letter}2:{letter}{total - 1})")
    cell.font, cell.border = Font(bold=True), BORDER
summary.cell(row=total + 2, column=2, value="Attente 20 (tarification claire) : vérifiée sur l'offre, sans scénario.").font = Font(italic=True, color="666666")

for ws in (guide, summary):
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

wb.properties.creator = "E-VOLUTION XP"
wb.properties.title = "Cahier de recette EGUITRA Finance et AxisPro Suite"
wb.save(OUT)
print(f"Écrit : {OUT} ({len(SCENARIOS)} scénarios)")
