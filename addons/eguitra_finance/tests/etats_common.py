"""Outils des tests des états au modèle : classeur de synthèse au format du modèle du client (jamais ses
chiffres), comptes et écritures d'essai."""

import io

from odoo import Command

M = 1_000_000

# Exercice N du classeur d'essai, en millions de GNF ; actif : (brut, amortissements et dépréciations, net).
# Comme dans le modèle du client : AI comprend AP, CB est porté en positif, les lignes « (-) » des flux aussi
# (ZB = FA - FB - FC - FD + FE). Ramenés à la présentation de la plateforme, les états sont cohérents :
# bilan équilibré, CJ = XI, totaux justes, ZH = BT - DT, ZA = BT - DT de N-1, report à nouveau justifié.
ACTIF = {
    "AD": (50, 10, 40), "AE": (0, 0, 0), "AF": (50, 10, 40), "AG": (0, 0, 0), "AH": (0, 0, 0),
    "AI": (1085, 135, 950), "AJ": (300, 0, 300), "AK": (0, 0, 0), "AL": (100, 20, 80), "AM": (60, 15, 45),
    "AN": (600, 100, 500), "AP": (25, 0, 25), "AQ": (10, 0, 10), "AR": (0, 0, 0), "AS": (10, 0, 10),
    "AZ": (1145, 145, 1000), "BA": (0, 0, 0), "BB": (40, 5, 35), "BG": (515, 20, 495), "BH": (0, 0, 0),
    "BI": (500, 20, 480), "BJ": (15, 0, 15), "BK": (555, 25, 530), "BQ": (0, 0, 0), "BR": (0, 0, 0),
    "BS": (120, 0, 120), "BT": (120, 0, 120), "BU": (0, 0, 0), "BZ": (1820, 170, 1650),
}
ACTIF_PREVIOUS = {"BS": 150, "BT": 150}
PASSIF = {
    "CA": 100, "CB": 10, "CD": 0, "CE": 0, "CF": 20, "CG": 30, "CH": 200, "CJ": 300, "CL": 0, "CM": 0, "CP": 640,
    "DA": 400, "DB": 0, "DC": 0, "DD": 400, "DF": 1040, "DH": 0, "DI": 0, "DJ": 350, "DK": 100, "DM": 60, "DN": 0,
    "DP": 510, "DQ": 0, "DR": 100, "DT": 100, "DV": 0, "DZ": 1650,
}
PASSIF_PREVIOUS = {"CF": 10, "CG": 10, "CH": 100, "CJ": 180, "DR": 70, "DT": 70}
RESULTAT = {
    "TA": 0, "RA": 0, "RB": 0, "XA": 0, "TB": 0, "TC": 2000, "TD": 0, "XB": 2000, "TE": 0, "TF": 0, "TG": 0,
    "TH": 0, "TI": 0, "RC": 0, "RD": 0, "RE": 400, "RF": 0, "RG": 100, "RH": 300, "RI": 50, "RJ": 0, "XC": 1150,
    "RK": 500, "XD": 650, "TJ": 0, "RL": 135, "XE": 515, "TK": 0, "TL": 0, "TM": 0, "RM": 50, "RN": 0, "XF": -50,
    "XG": 465, "TN": 0, "TO": 0, "RO": 0, "RP": 0, "XH": 0, "RQ": 0, "RS": 165, "XI": 300,
}
FLUX = {
    "ZA": 80, "FA": 435, "FB": 0, "FC": 10, "FD": 100, "FE": -5, "ZB": 320, "FF": 0, "FG": 300, "FH": 0,
    "FI": 0, "FJ": 0, "ZC": -300, "FK": 0, "FL": 0, "FM": 0, "FN": 50, "ZD": -50, "FO": 0, "FP": 0, "FQ": 30,
    "ZE": -30, "ZF": -80, "ZG": -60, "ZH": 20,
}
LETTERS = {"ZA": "A", "ZB": "B", "ZC": "C", "ZD": "D", "ZE": "E", "ZF": "F", "ZG": "G", "ZH": "H"}


def client_workbook(year, errors=None):
    """Classeur de synthèse au format du modèle du client : feuilles, lignes d'en-tête « REF », intitulés et
    fusions de colonnes, ligne AP sans brut ni amortissement, intertitres et ligne du BF sans REF dans les flux.
    errors : {code: montant net N en millions} portés à la place du bon montant."""
    import openpyxl

    errors = errors or {}

    def value(code, amount):
        return errors.get(code, amount) * M

    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "Bilan"
    ws["A2"] = "Modèle du Bilan du Système Normal OHADA"
    ws["A3"] = "BILAN ACTIF"
    ws["A4"] = "Désignation entité : SOCIÉTÉ D'ESSAI SARLU"
    ws["A5"] = "Numéro d’identification : 000000000"
    ws["B5"] = "Durée (en mois) 12"
    ws["B6"] = f"BILAN AU 31 DECEMBRE {year}"
    for col, text in zip("ABCDGH", ("REF", "ACTIF", "Note", "EXERCICE AU         31/12/N",
                                    "EXERCICE AU         31/12/N-1", "Formules"), strict=True):
        ws[f"{col}8"] = text
    ws.merge_cells("D8:F9")
    ws.merge_cells("G8:G9")
    for col, text in zip("DEFG", ("BRUT", "AMORT et DEPREC.", "NET", "NET"), strict=True):
        ws[f"{col}10"] = text
    row = 11
    for code, (gross, depreciation, net) in ACTIF.items():
        ws.cell(row, 1, code)
        ws.cell(row, 2, f"Poste {code}")
        if code != "AP":
            ws.cell(row, 4, gross * M)
            ws.cell(row, 5, depreciation * M)
        ws.cell(row, 6, value(code, net))
        if code in ACTIF_PREVIOUS:
            ws.cell(row, 7, ACTIF_PREVIOUS[code] * M)
        row += 1
    row += 1
    ws.cell(row, 1, "BILAN PASSIF")
    ws.cell(row + 1, 1, "Désignation entité : SOCIÉTÉ D'ESSAI SARLU                Exercice clos le 31-")
    ws.cell(row + 2, 1, "Numéro d’identification : 000000000                       Durée (en mois")
    ws.cell(row + 3, 2, "BILAN AU 31 DECEMBRE N")
    row += 4
    for col, text in enumerate(("REF", "PASSIF", "Note", f"EXERCICE AU  31/12/{year}", f"EXERCICE AU   31/12/{year - 1}",
                                "Formules"), start=1):
        ws.cell(row, col, text)
    ws.cell(row + 1, 4, "NET")
    ws.cell(row + 1, 5, "NET")
    row += 2
    for code, amount in PASSIF.items():
        ws.cell(row, 1, code)
        ws.cell(row, 2, "Apporteurs capital non appelé          (-)" if code == "CB" else f"Poste {code}")
        ws.cell(row, 4, value(code, amount))
        if code in PASSIF_PREVIOUS:
            ws.cell(row, 5, PASSIF_PREVIOUS[code] * M)
        row += 1

    ws = book.create_sheet("Compte de Résultat")
    ws["A1"] = "Modèle de Compte de résultat du Système Normal OHADA"
    ws["A2"] = "Désignation entité : SOCIÉTÉ D'ESSAI SARLU"
    ws["B3"] = "COMPTE DE RESULTAT AU 31 DECEMBRE N"
    for col, text in zip("ABDEF", ("REF", "LIBELLES", "NOTE", "EXERCICE", "EXERCICE"), strict=True):
        ws[f"{col}4"] = text
    ws.merge_cells("C4:C7")
    ws["E5"], ws["F5"] = "AU         31/12/N", "AU   31/12/N-1"
    ws["E7"], ws["F7"] = "NET", "NET"
    for row, (code, amount) in enumerate(RESULTAT.items(), start=8):
        ws.cell(row, 1, code)
        ws.cell(row, 2, f"Rubrique {code}")
        ws.cell(row, 3, {"T": "+", "R": "-"}.get(code[0]))
        ws.cell(row, 5, value(code, amount))

    ws = book.create_sheet("Tableau des  Flux de trésorerie")
    ws["A1"] = "Modèle du Tableau des flux de trésorerie du Système Normal OHADA"
    for col, text in zip("ABDE", ("REF", "LIBELLES", f"EXERCICE {year}", f"EXERCICE {year - 1}"), strict=True):
        ws[f"{col}3"] = text
    row = 4
    for code, amount in FLUX.items():
        if code == "ZB":
            ws.cell(row, 2, "Variation du BF lié aux activités opérationnelles (FB+FC+FD+FE)")
            ws.cell(row, 4, (-FLUX["FB"] - FLUX["FC"] - FLUX["FD"] + FLUX["FE"]) * M)
            row += 1
        if code in ("FA", "FF", "FK", "FO"):
            ws.cell(row, 2, "Intertitre")
            row += 1
        ws.cell(row, 1, code)
        ws.cell(row, 2, f"Ligne {code}")
        ws.cell(row, 3, LETTERS.get(code))
        ws.cell(row, 4, value(code, amount))
        row += 1
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def plan_account(env, company, prefix, kind):
    """Compte du plan de la société commençant par prefix ; créé s'il manque au plan d'essai."""
    Account = env["account.account"]
    found = Account.search([("code", "=like", prefix + "%"), ("company_ids", "in", company.ids)], order="code", limit=1)
    return found or Account.create({"code": prefix.ljust(6, "0"), "name": f"Compte {prefix} (essai)", "account_type": kind,
                                    "company_ids": [(6, 0, company.ids)]})


def general_journal(env, company, code=None):
    """Journal des opérations diverses (hors à-nouveaux), ou le journal AN (créé s'il manque)."""
    Journal = env["account.journal"]
    if code == "AN":
        return Journal.search([("company_id", "=", company.id), ("code", "=", "AN")], limit=1) or Journal.create(
            {"name": "À-nouveaux", "code": "AN", "type": "general", "company_id": company.id})
    return Journal.search([("company_id", "=", company.id), ("type", "=", "general"), ("code", "!=", "AN")],
                          order="sequence, id", limit=1)


def post_entry(env, journal, day, lines):
    """Écriture validée : lines = [(compte, débit, crédit[, tiers])]."""
    move = env["account.move"].create({
        "move_type": "entry", "journal_id": journal.id, "date": day,
        "line_ids": [Command.create({"name": "Essai", "account_id": line[0].id, "debit": line[1], "credit": line[2],
                                     "partner_id": line[3].id if len(line) > 3 else False}) for line in lines],
    })
    move.with_context(eguitra_approval_granted=True).action_post()
    return move
