import io
from datetime import date

from dateutil.relativedelta import relativedelta

from odoo.exceptions import UserError
from odoo.tests import TransactionCase, tagged

from odoo.addons.eguitra_api.files import XLSX
from odoo.addons.eguitra_api.tests.common import ApiCase

from ..controllers.etats_modele import SHEETS, TREASURY_CHECK
from ..models.statement_comparative import ComparativeExists, parse_workbook
from ..models.syscohada import (
    BILAN_ACTIF_ORDER,
    BILAN_ACTIF_TOTALS,
    BILAN_PASSIF_ORDER,
    CR_ORDER,
    TFT_LETTERS,
    TFT_ROWS,
    WORKING_CAPITAL_LABEL,
)
from .common import FinanceDataMixin
from .etats_common import M, client_workbook, general_journal, plan_account, post_entry
from .exercices_common import PREVIOUS, TwoYearsMixin

TFT_CODES = [code for code, _label, _total in TFT_ROWS]
# Comparatifs importés sur un exercice ancien, sans aucune écriture dans la plateforme.
CLOSING = date(2015, 12, 31)
NEXT_START, NEXT_END = date(2016, 1, 1), date(2016, 12, 31)
BAD = {"DZ": 1640, "XI": 290, "AD": 45, "ZH": 25, "CH": 210}


@tagged("post_install", "-at_install")
class TestBilanDetaille(FinanceDataMixin, TransactionCase):
    """Bilan au modèle du système normal : postes de l'actif détaillés (AE à AH, AJ à AN, AR, AS), brut,
    amortissements et dépréciations et net de chaque poste (immobilisations 28 et 29, stocks 39, créances 49,
    trésorerie 59), notes annexes ; signe et notes du compte de résultat ; lettres du tableau des flux."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.setup_finance_data()
        env, company, day = cls.env, cls.company, cls.today

        def acc(prefix, kind):
            return plan_account(env, company, prefix, kind)

        cls.fy_start = env["eguitra.finance.metrics"].fiscal_year_bounds(company, day)[0]
        various = general_journal(env, company)
        bank, capital = cls.bank.default_account_id, acc("101", "equity")
        # Apport en numéraire ; camion (245) et terrain (22) payés par la banque.
        post_entry(env, various, day, [(bank, 1000 * M, 0), (capital, 0, 1000 * M)])
        post_entry(env, various, day, [(acc("245", "asset_fixed"), 600 * M, 0), (acc("222", "asset_fixed"), 300 * M, 0),
                                       (bank, 0, 900 * M)])
        # Apports en nature : logiciel (213), bâtiment (231), marchandises (31).
        post_entry(env, various, day, [(acc("2131", "asset_fixed"), 50 * M, 0), (acc("231", "asset_fixed"), 200 * M, 0),
                                       (acc("31", "asset_current"), 40 * M, 0), (capital, 0, 290 * M)])
        # Amortissements (2845, 2813, 2831) et dépréciations des stocks (39) et des clients (491).
        post_entry(env, various, day, [
            (acc("6813", "expense_depreciation"), 130 * M, 0), (cls.acc_purchase, 25 * M, 0),
            (acc("2845", "asset_fixed"), 0, 100 * M), (acc("2813", "asset_fixed"), 0, 10 * M),
            (acc("2831", "asset_fixed"), 0, 20 * M), (acc("39", "asset_current"), 0, 5 * M),
            (acc("491", "asset_current"), 0, 20 * M),
        ])
        receivable = cls.customer.property_account_receivable_id
        post_entry(env, various, day, [(receivable, 500 * M, 0, cls.customer), (cls.acc_services, 0, 500 * M)])

    def test_postes_brut_amortissements_et_net(self):
        sheet = self.env["eguitra.syscohada"].balance_sheet(self.company, self.today)
        self.assertEqual([row["code"] for row in sheet["actif"]], BILAN_ACTIF_ORDER, "Ordre du modèle")
        self.assertEqual([row["code"] for row in sheet["passif"]], BILAN_PASSIF_ORDER)
        rows = {row["code"]: row for row in sheet["actif"]}
        expected = {
            "AN": (600, 100, 500), "AJ": (300, 0, 300), "AF": (50, 10, 40), "AK": (200, 20, 180), "AL": (0, 0, 0),
            "AD": (50, 10, 40), "AI": (1100, 120, 980), "AP": (0, 0, 0), "AQ": (0, 0, 0), "AZ": (1150, 130, 1020),
            "BB": (40, 5, 35), "BI": (500, 20, 480), "BG": (500, 20, 480), "BK": (540, 25, 515), "BS": (100, 0, 100),
            "BT": (100, 0, 100), "BZ": (1790, 155, 1635),
        }
        for code, (gross, depreciation, net) in expected.items():
            row = rows[code]
            self.assertAlmostEqual(row["gross"], gross * M, 0, msg=f"{code} brut")
            self.assertAlmostEqual(row["depreciation"], depreciation * M, 0, msg=f"{code} amortissements et dépréciations")
            self.assertAlmostEqual(row["amount"], net * M, 0, msg=f"{code} net")
        for code, _label, formula in BILAN_ACTIF_TOTALS:
            for key in ("gross", "depreciation", "amount"):
                self.assertAlmostEqual(rows[code][key], sum(rows[c][key] * s for c, s in formula), 2, msg=(code, key))
        for row in sheet["actif"]:
            self.assertAlmostEqual(row["amount"], row["gross"] - row["depreciation"], 2, msg=row["code"])
        self.assertEqual(sheet["difference"], 0, "Bilan équilibré")
        self.assertEqual(sheet["unmapped_accounts"], {}, "Aucun compte hors des postes")
        self.assertEqual((rows["AD"]["note"], rows["AQ"]["note"], rows["BI"]["note"], rows["AN"]["note"]), ("3", "4", "7", None))
        self.assertTrue(rows["AI"]["total"] and not rows["AI"]["detail"], "AI : rubrique, somme de AJ à AN")
        self.assertTrue(rows["AN"]["detail"] and not rows["AN"]["total"])
        self.assertFalse(rows["AP"]["total"] or rows["AP"]["detail"], "AP : rubrique propre (présentation officielle)")
        passif = {row["code"]: row for row in sheet["passif"]}
        self.assertEqual((passif["CE"]["note"], passif["CJ"]["note"], passif["DJ"]["note"]), ("3e", None, "17"))
        self.assertAlmostEqual(passif["CJ"]["amount"], 345 * M, 0, msg="Services 500, dotations 130, dépréciations 25")

    def test_compte_de_resultat_signes_et_notes(self):
        statement = self.env["eguitra.syscohada"].income_statement(self.company, self.fy_start, self.today)
        lines = {line["code"]: line for line in statement["lines"]}
        self.assertEqual((lines["TC"]["sign"], lines["TC"]["note"]), ("+", "21"))
        self.assertEqual((lines["RE"]["sign"], lines["RE"]["note"]), ("-", "22"))
        self.assertEqual((lines["RL"]["sign"], lines["RL"]["note"]), ("-", "3C&28"))
        self.assertEqual((lines["XB"]["sign"], lines["XB"]["note"]), (None, None), "Solde intermédiaire sans signe")
        self.assertIsNone(lines["RS"]["note"])
        for line in statement["lines"]:
            self.assertEqual(line["sign"] is None, line["total"], line["code"])
        self.assertAlmostEqual(lines["RL"]["amount"], 130 * M, 0, msg="Montants positifs, le signe dit le sens")

    def test_tableau_des_flux_lettres_et_bf(self):
        flow = self.env["eguitra.syscohada"].cash_flow(self.company, self.fy_start, self.today)
        lines = {line["code"]: line for line in flow["lines"]}
        self.assertEqual({code: line["letter"] for code, line in lines.items() if line["letter"]}, TFT_LETTERS)
        self.assertAlmostEqual(flow["bf"], sum(lines[c]["amount"] for c in ("FB", "FC", "FD", "FE")), 2)
        self.assertAlmostEqual(lines["ZH"]["amount"], lines["ZG"]["amount"] + lines["ZA"]["amount"], 2)
        self.assertAlmostEqual(flow["treasury_end"], 100 * M, 0, msg="Contrôle de ZH : trésorerie actif - passif")
        self.assertAlmostEqual(flow["difference"], 0, 0, msg="ZH rejoint la trésorerie du bilan")


@tagged("post_install", "-at_install")
class TestColonneN1(TwoYearsMixin, TransactionCase):
    """Colonne N-1 lue dans la plateforme quand l'exercice précédent y a ses écritures : tableau des flux sur la
    même période un an plus tôt, bilan à la clôture précédente ; des comparatifs importés n'y servent pas."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.setup_two_years()

    def test_tableau_des_flux_n1(self):
        Syscohada = self.env["eguitra.syscohada"]
        flows = Syscohada.cash_flow_compared(self.company, self.fy_start, self.fy_end)
        previous = Syscohada.cash_flow(self.company, self.previous_start, self.previous_end)
        expected = {line["code"]: line["amount"] for line in previous["lines"]}
        self.assertAlmostEqual(expected["FA"], PREVIOUS, 0, msg="CAFG N-1 : 1 000 M de services, 200 M de charges")
        for line in flows["lines"]:
            self.assertAlmostEqual(line["previous"], expected[line["code"]], 2, msg=line["code"])
        meta = flows["previous"]
        self.assertEqual((meta["source"], meta["date_from"], meta["date_to"]), ("platform", self.previous_start, self.previous_end))
        self.assertAlmostEqual(meta["bf"], sum(expected[c] for c in ("FB", "FC", "FD", "FE")), 2)
        self.assertAlmostEqual(meta["treasury_end"], previous["treasury_end"], 2)
        self.assertAlmostEqual(meta["difference"], previous["difference"], 2)

    def test_bilan_n1_a_la_cloture_precedente(self):
        Syscohada = self.env["eguitra.syscohada"]
        sheet = Syscohada.balance_sheet_compared(self.company, self.today)
        closing = Syscohada.balance_sheet(self.company, self.previous_end)
        expected = {row["code"]: row["amount"] for row in closing["actif"] + closing["passif"]}
        self.assertEqual(sheet["previous"]["date_to"], self.previous_end)
        for row in sheet["actif"] + sheet["passif"]:
            self.assertAlmostEqual(row["previous"], expected[row["code"]], 2, msg=row["code"])
        self.assertAlmostEqual(expected["CJ"], PREVIOUS, 0, msg="Résultat N-1 au bilan de clôture")

    def test_comparatifs_ignores_quand_l_exercice_precedent_est_dans_la_plateforme(self):
        self.env["eguitra.statement.comparative"].import_workbook(
            self.company, self.previous_end, client_workbook(self.previous_end.year), "etats.xlsx")
        Syscohada = self.env["eguitra.syscohada"]
        income = Syscohada.previous_income(self.company, self.fy_start, self.fy_end)
        self.assertEqual(income["source"], "platform")
        self.assertFalse(income["import_available"])
        self.assertAlmostEqual(income["amounts"]["TC"], 1_000_000_000, 0, msg="Chiffre de la plateforme, pas du classeur")
        self.assertEqual(Syscohada.balance_sheet_compared(self.company, self.today)["previous"]["source"], "platform")


@tagged("post_install", "-at_install")
class TestComparatifs(FinanceDataMixin, TransactionCase):
    """Comparatifs du premier exercice : lecture du modèle par les codes REF et les intitulés des colonnes,
    conversions vers la présentation officielle, contrôles sans blocage, emploi comme colonne N-1, bilan
    d'ouverture face aux à-nouveaux."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.setup_finance_data()
        cls.Comparative = cls.env["eguitra.statement.comparative"]
        cls.Syscohada = cls.env["eguitra.syscohada"]

    def test_lecture_du_modele_et_conversions(self):
        parsed = parse_workbook(client_workbook(2015))
        v = parsed["values"]
        self.assertEqual(parsed["years"], [2014, 2015])
        self.assertEqual((parsed["entity"], parsed["vat"]), ("SOCIÉTÉ D'ESSAI SARLU", "000000000"))
        self.assertEqual((v["AN"]["gross"], v["AN"]["depreciation"], v["AN"]["amount"]), (600 * M, 100 * M, 500 * M))
        self.assertEqual((v["AI"]["gross"], v["AI"]["depreciation"], v["AI"]["amount"]), (1060 * M, 135 * M, 925 * M),
                         "AI ramené à AJ à AN : AP est une rubrique propre")
        self.assertEqual((v["AP"]["gross"], v["AP"]["amount"]), (25 * M, 25 * M), "Net seul : brut = net")
        self.assertEqual(v["CB"]["amount"], -10 * M, "Capital non appelé en déduction")
        self.assertEqual((v["FC"]["amount"], v["FD"]["amount"], v["FG"]["amount"], v["FN"]["amount"], v["FQ"]["amount"]),
                         (-10 * M, -100 * M, -300 * M, -50 * M, -30 * M), "Lignes « (-) » retenues en négatif")
        self.assertEqual(v["FE"]["amount"], -5 * M, "Ligne « (+) » inchangée")
        self.assertEqual((v["BT"]["previous"], v["CH"]["previous"], v["XI"]["amount"]), (150 * M, 100 * M, 300 * M))
        notes = " ".join(parsed["notes"])
        for text in ("AI comprenait AP", "CB (apporteurs", "ZB :", "ZC :", "ZD :", "ZE :"):
            self.assertIn(text, notes)

    def test_controles_affiches_sans_bloquer(self):
        record = self.Comparative.import_workbook(self.company, CLOSING, client_workbook(2015), "etats_2015.xlsx")
        self.assertEqual((record.date_from, record.date_to, record.source), (date(2015, 1, 1), CLOSING, "etats_2015.xlsx"))
        self.assertTrue(record.has_previous)
        self.assertEqual(record.imported_by, self.env.user)
        checks = {c["code"]: c for c in record.controls()}
        self.assertEqual(set(checks), {"bilan_equilibre", "resultat", "totaux", "net_actif", "tresorerie_fin",
                                       "tresorerie_debut", "report_a_nouveau"})
        self.assertTrue(all(c["ok"] for c in checks.values()), checks)
        summary = record.summary()
        self.assertEqual((summary["conform"], summary["total"]), (7, 7))

        with self.assertRaises(ComparativeExists):
            self.Comparative.import_workbook(self.company, CLOSING, client_workbook(2015))
        replaced = self.Comparative.import_workbook(self.company, CLOSING, client_workbook(2015, errors=BAD), "faux.xlsx",
                                                    replace=True)
        self.assertEqual(replaced, record, "Remplacement en place")
        self.assertEqual(record.source, "faux.xlsx")
        checks = {c["code"]: c for c in record.controls()}
        self.assertEqual((checks["bilan_equilibre"]["ok"], checks["bilan_equilibre"]["gap"]), (False, -10 * M), "BZ = DZ")
        self.assertEqual((checks["resultat"]["ok"], checks["resultat"]["gap"]), (False, 10 * M), "CJ = XI")
        self.assertFalse(checks["totaux"]["ok"])
        failing = {(item["code"], item["field"]) for item in checks["totaux"]["items"]}
        self.assertLessEqual({("AD", "amount"), ("AZ", "amount"), ("CP", "amount"), ("DZ", "amount"), ("XI", "amount"),
                              ("ZH", "amount")}, failing)
        self.assertEqual([item["code"] for item in checks["net_actif"]["items"]], ["AD"], "Net = brut - amortissements")
        self.assertEqual((checks["tresorerie_fin"]["ok"], checks["tresorerie_fin"]["gap"]), (False, 5 * M), "ZH = BT - DT")
        self.assertTrue(checks["tresorerie_debut"]["ok"])
        self.assertEqual((checks["report_a_nouveau"]["ok"], checks["report_a_nouveau"]["gap"]), (False, 10 * M))
        self.assertEqual(self.Comparative.search_count([("company_id", "=", self.company.id), ("date_to", "=", CLOSING)]), 1)

    def test_refus_de_l_import(self):
        with self.assertRaises(UserError, msg="Pas une date de clôture"):
            self.Comparative.import_workbook(self.company, date(2015, 6, 30), client_workbook(2015))
        with self.assertRaises(UserError, msg="Classeur d'un autre exercice"):
            self.Comparative.import_workbook(self.company, CLOSING, client_workbook(2014))
        import openpyxl

        book, out = openpyxl.Workbook(), io.BytesIO()
        book.active["A1"] = "Sans colonne REF"
        book.save(out)
        with self.assertRaises(UserError, msg="Classeur hors modèle"):
            self.Comparative.import_workbook(self.company, CLOSING, out.getvalue())
        self.assertFalse(self.Comparative.search([("company_id", "=", self.company.id)]))

    def test_comparatifs_en_colonne_n1(self):
        record = self.Comparative.import_workbook(self.company, CLOSING, client_workbook(2015))
        sheet = self.Syscohada.balance_sheet_compared(self.company, NEXT_END)
        self.assertEqual((sheet["previous"]["source"], sheet["previous"]["date_to"]), ("import", CLOSING))
        self.assertEqual([i["id"] for i in sheet["previous"]["imports"]], record.ids)
        actif = {row["code"]: row["previous"] for row in sheet["actif"]}
        passif = {row["code"]: row["previous"] for row in sheet["passif"]}
        self.assertEqual((actif["BI"], actif["AN"], actif["AI"], actif["AP"], actif["BZ"]),
                         (480 * M, 500 * M, 925 * M, 25 * M, 1650 * M))
        self.assertEqual((passif["CB"], passif["CP"], passif["DZ"]), (-10 * M, 640 * M, 1650 * M))
        self.assertEqual(sheet["previous"]["difference"], 0)

        income = self.Syscohada.add_previous_income(
            self.Syscohada.income_statement(self.company, NEXT_START, NEXT_END), self.company, NEXT_START, NEXT_END)
        lines = {line["code"]: line["previous"] for line in income["lines"]}
        self.assertEqual((lines["TC"], lines["RS"], lines["XI"]), (2000 * M, 165 * M, 300 * M))
        self.assertEqual(income["previous"]["source"], "import")

        flows = self.Syscohada.cash_flow_compared(self.company, NEXT_START, NEXT_END)
        lines = {line["code"]: line["previous"] for line in flows["lines"]}
        self.assertEqual((lines["ZA"], lines["FG"], lines["ZB"], lines["ZH"]), (80 * M, -300 * M, 320 * M, 20 * M))
        previous = flows["previous"]
        self.assertEqual((previous["bf"], previous["treasury_end"], previous["difference"]), (-115 * M, 20 * M, 0))

        # Période partielle : les comparatifs annuels ne s'emploient pas, l'écran le signale.
        partial = self.Syscohada.previous_income(self.company, NEXT_START, date(2016, 6, 30))
        self.assertEqual((partial["source"], partial["import_available"]), ("platform", True))
        self.assertAlmostEqual(partial["amounts"]["TC"], 0, 0)

    def test_bilan_d_ouverture(self):
        self.Comparative.import_workbook(self.company, CLOSING, client_workbook(2015))
        acc = lambda prefix, kind: plan_account(self.env, self.company, prefix, kind)  # noqa: E731
        post_entry(self.env, general_journal(self.env, self.company, "AN"), NEXT_START, [
            (self.bank.default_account_id, 120 * M, 0), (acc("245", "asset_fixed"), 600 * M, 0),
            (acc("2845", "asset_fixed"), 0, 100 * M), (acc("101", "equity"), 0, 620 * M),
        ])
        check = self.Comparative.opening_check(self.company, CLOSING)
        self.assertTrue(check["available"])
        self.assertEqual((check["opening_date"], check["accounts"]), (NEXT_START, 4))
        rows = {row["code"]: row for row in check["rows"]}
        self.assertTrue(rows["AN"]["ok"], rows["AN"])
        self.assertEqual(rows["AN"]["opening"], {"gross": 600 * M, "depreciation": 100 * M, "amount": 500 * M})
        self.assertTrue(rows["BS"]["ok"], "Banque reprise au solde du bilan importé")
        self.assertFalse(rows["CA"]["ok"])
        self.assertEqual(rows["CA"]["gaps"]["amount"], 520 * M, "Capital repris : 620 M au lieu de 100 M")
        self.assertEqual(check["gaps"], sum(1 for row in check["rows"] if not row["ok"]))
        self.assertGreater(check["gaps"], 2)
        self.assertFalse(self.Comparative.opening_check(self.company, date(2014, 12, 31))["available"])


@tagged("post_install", "-at_install")
class TestEtatsModeleApi(FinanceDataMixin, ApiCase):
    """Export au modèle (trois feuilles, en-têtes, codes, totaux en gras, valeurs sans formule), import des
    comparatifs par l'administration avec remplacement confirmé, refus par profil."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.setup_finance_data()
        cls._create_user("compta.modele", "finance")
        cls._create_user("admin.modele", "admin")
        cls._create_user("dg.modele", "direction")
        cls._create_user("parc.modele", "exploitation")
        cls.make_invoice("out_invoice", 200_000_000)
        cls.make_invoice("in_invoice", 50_000_000)
        cls.fy_start, cls.fy_end = cls.env["eguitra.finance.metrics"].fiscal_year_bounds(cls.company, cls.today)

    def _token(self, login):
        return self.open_session(login)["access_token"]

    def _download(self, token, path):
        return self.url_open("/api/v1" + path, headers={"Authorization": f"Bearer {token}", "X-Eguitra-App": "test"},
                             timeout=60)

    def _upload(self, token, content, **form):
        response = self.url_open("/api/v1/finance/comparatives/import", data=form,
                                 files={"file": ("etats_2015.xlsx", io.BytesIO(content), XLSX)},
                                 headers={"Authorization": f"Bearer {token}", "X-Eguitra-App": "test"},
                                 method="POST", timeout=60)
        return response.status_code, response.json()

    def test_export_au_modele(self):
        import openpyxl

        self.env.flush_all()
        token = self._token("compta.modele")  # une session par compte : la seconde demanderait le code TOTP
        response = self._download(token, f"/finance/statements/export?kind=modele_syscohada&date_to={self.fy_end.isoformat()}")
        self.assertEqual(response.status_code, 200, response.text[:300])
        book = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(book.sheetnames, list(SHEETS), "Feuilles nommées comme dans le modèle")
        cells = {name: [list(row) for row in book[name].iter_rows()] for name in SHEETS}
        values = {name: [[c.value for c in row] for row in rows] for name, rows in cells.items()}
        formulas = [c.coordinate for rows in cells.values() for row in rows for c in row
                    if isinstance(c.value, str) and c.value.startswith("=")]
        self.assertEqual(formulas, [], "Des valeurs, jamais de formules")
        previous_end = self.fy_start - relativedelta(days=1)
        texts = {str(v) for rows in values.values() for row in rows for v in row if v not in (None, "")}
        for text in (f"Désignation entité : {self.company.name}", f"Exercice clos le {self.fy_end:%d-%m-%Y}",
                     "Durée (en mois) : 12", f"BILAN AU 31 DÉCEMBRE {self.fy_end.year}",
                     f"COMPTE DE RÉSULTAT AU 31 DÉCEMBRE {self.fy_end.year}", WORKING_CAPITAL_LABEL, TREASURY_CHECK):
            self.assertIn(text, texts)
        self.assertTrue(any(t.startswith("Numéro d'identification : ") for t in texts))

        bilan, rows = values["Bilan"], cells["Bilan"]
        heads = [i for i, row in enumerate(bilan) if row[0] == "REF"]
        self.assertEqual(len(heads), 2, "Actif puis passif")
        actif, passif = heads
        self.assertEqual(bilan[actif][:4], ["REF", "ACTIF", "Note", f"EXERCICE AU {self.fy_end:%d/%m/%Y}"])
        self.assertEqual(bilan[actif][6], f"EXERCICE AU {previous_end:%d/%m/%Y}")
        self.assertEqual(bilan[actif + 1][3:7], ["BRUT", "AMORT. ET DÉPRÉC.", "NET", "NET"])
        self.assertEqual([row[0] for row in bilan[actif + 2:actif + 2 + len(BILAN_ACTIF_ORDER)]], BILAN_ACTIF_ORDER)
        self.assertEqual(bilan[passif][:5], ["REF", "PASSIF", "Note", f"EXERCICE AU {self.fy_end:%d/%m/%Y}",
                                             f"EXERCICE AU {previous_end:%d/%m/%Y}"])
        self.assertEqual([row[0] for row in bilan[passif + 2:passif + 2 + len(BILAN_PASSIF_ORDER)]], BILAN_PASSIF_ORDER)
        line = {row[0]: (i, row) for i, row in enumerate(bilan) if row[0]}
        for code in ("AZ", "BZ", "CP", "DZ"):
            self.assertTrue(rows[line[code][0]][1].font.b, f"{code} en gras")
        self.assertFalse(rows[line["BI"][0]][1].font.b, "Poste en maigre")
        sheet = self.env["eguitra.syscohada"].balance_sheet(self.company, self.fy_end)
        bi = next(row for row in sheet["actif"] if row["code"] == "BI")
        self.assertEqual(line["BI"][1][2], "7", "Note")
        self.assertAlmostEqual(line["BI"][1][3], bi["gross"], 2)
        self.assertAlmostEqual(line["BI"][1][5], bi["amount"], 2)
        dj = next(row for row in sheet["passif"] if row["code"] == "DJ")
        self.assertAlmostEqual(line["DJ"][1][3], dj["amount"], 2)

        resultat = values["Compte de Résultat"]
        head = next(i for i, row in enumerate(resultat) if row[0] == "REF")
        self.assertEqual([resultat[head][i] for i in (0, 1, 3)], ["REF", "LIBELLÉS", "NOTE"])
        self.assertIn(resultat[head][2], (None, ""), "Colonne du signe")
        self.assertEqual([row[0] for row in resultat[head + 2:head + 2 + len(CR_ORDER)]], CR_ORDER)
        rubriques = {row[0]: row for row in resultat[head + 2:]}
        self.assertEqual(rubriques["TC"][2:4], ["+", "21"])
        self.assertEqual(rubriques["RL"][2:4], ["-", "3C&28"])
        self.assertIn(rubriques["XB"][2], (None, ""))
        self.assertAlmostEqual(rubriques["TC"][4], 200_000_000, 0)

        flux = values["Tableau des  Flux de trésorerie"]
        head = next(i for i, row in enumerate(flux) if row[0] == "REF")
        self.assertEqual(flux[head][:2], ["REF", "LIBELLÉS"])
        self.assertEqual(flux[head][3:5], [f"EXERCICE {self.fy_end.year}", f"EXERCICE {self.fy_end.year - 1}"])
        coded = [row for row in flux[head + 1:] if row[0]]
        self.assertEqual([row[0] for row in coded], TFT_CODES)
        self.assertEqual({row[0]: row[2] for row in coded if row[2]}, TFT_LETTERS)

        # Le classeur produit se relit par l'import : mêmes montants, aucune conversion.
        parsed = parse_workbook(response.content)
        self.assertEqual(parsed["notes"], [])
        self.assertAlmostEqual(parsed["values"]["BI"]["amount"], bi["amount"], 2)
        self.assertAlmostEqual(parsed["values"]["BI"]["depreciation"], bi["depreciation"], 2)
        self.assertAlmostEqual(parsed["values"]["DJ"]["amount"], dj["amount"], 2)

        status, body = self.call("GET", "/finance/statements/export?kind=modele_syscohada&format=csv", token=token)
        self.assertEqual(status, 400, "Un classeur XLSX seulement")

    def test_import_et_remplacement_confirme(self):
        admin = self._token("admin.modele")
        status, body = self._upload(admin, client_workbook(2015), date_to="2015-12-31")
        self.assertEqual(status, 200, body)
        self.assertEqual((body["date_to"], body["conform"], body["total"], body["replaced"]), ("2015-12-31", 7, 7, False))
        self.assertTrue(any("AI comprenait AP" in note for note in body["notes"]))
        first = body["id"]

        status, body = self._upload(admin, client_workbook(2015, errors=BAD), date_to="2015-12-31")
        self.assertEqual((status, body["error"]), (409, "conflict"), "Remplacement à confirmer")
        self.assertEqual(body["existing"]["id"], first)
        record = self.env["eguitra.statement.comparative"].browse(first)
        self.assertTrue(all(c["ok"] for c in record.controls()), "Import précédent intact")

        status, body = self._upload(admin, client_workbook(2015, errors=BAD), date_to="2015-12-31", replace="1")
        self.assertEqual(status, 200, body)
        self.assertEqual((body["id"], body["replaced"]), (first, True))
        checks = {c["code"]: c for c in body["checks"]}
        self.assertFalse(checks["bilan_equilibre"]["ok"], "Contrôles affichés, import enregistré")
        self.assertEqual(self.env["eguitra.statement.comparative"].search_count([("company_id", "=", self.company.id)]), 1)
        for model in ("eguitra.statement.comparative", "eguitra.statement.comparative.line"):
            self.assertTrue(self.env["auditlog.rule"].sudo().search_count(
                [("model_id.model", "=", model), ("log_write", "=", True), ("log_unlink", "=", True)]), model)

        status, body = self._upload(admin, b"pas un classeur", date_to="2015-12-31")
        self.assertEqual(status, 422, body)

        finance = self._token("compta.modele")
        status, body = self.call("GET", "/finance/comparatives", token=finance)
        self.assertEqual(status, 200, body)
        self.assertEqual([item["id"] for item in body["items"]], [first])
        status, body = self.call("GET", "/finance/comparatives/opening?date_to=2015-12-31", token=finance)
        self.assertEqual(status, 200, body)
        self.assertTrue(body["available"])
        self.assertEqual(body["accounts"], 0, "Aucun à-nouveau en 2016")

    def test_droits_par_profil(self):
        parc = self._token("parc.modele")
        for path in ("/finance/statements/export?kind=modele_syscohada", "/finance/comparatives",
                     "/finance/comparatives/opening", "/finance/statements/balance-sheet"):
            status, body = self.call("GET", path, token=parc)
            self.assertEqual(status, 403, (path, body))
        status, body = self._upload(parc, client_workbook(2015), date_to="2015-12-31")
        self.assertEqual(status, 403, body)
        for login in ("compta.modele", "dg.modele"):
            token = self._token(login)
            response = self._download(token, "/finance/statements/export?kind=modele_syscohada")
            self.assertEqual(response.status_code, 200, (login, response.text[:200]))
            status, body = self.call("GET", "/finance/comparatives", token=token)
            self.assertEqual(status, 200, (login, body))
            status, body = self._upload(token, client_workbook(2015), date_to="2015-12-31")
            self.assertEqual(status, 403, (login, "Import réservé au profil de la validation de la reprise"))
        self.assertFalse(self.env["eguitra.statement.comparative"].search([("company_id", "=", self.company.id)]))
