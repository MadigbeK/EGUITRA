import csv
import io
from datetime import timedelta

from odoo import Command
from odoo.tests import tagged

from odoo.addons.eguitra_api.tests.common import ApiCase

from ..controllers.statements_export import CASH_STATUS
from .common import FinanceDataMixin


@tagged("post_install", "-at_install")
class TestStatementsExport(FinanceDataMixin, ApiCase):
    """Exports des états : un onglet par état, balance avec classes et soldes, grand livre avec
    activité et totaux, compte de résultat avec le mois et la part du chiffre d'affaires."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.setup_finance_data()
        cls._create_user("hawa", "finance")

    def setUp(self):
        super().setUp()
        self.token = self.open_session("hawa")["access_token"]
        self.make_invoice("out_invoice", 200_000_000)
        self.make_invoice("in_invoice", 50_000_000)

    def _get(self, path):
        response = self.url_open("/api/v1" + path, headers={"Authorization": f"Bearer {self.token}"}, timeout=60)
        self.assertEqual(response.status_code, 200, response.text[:300])
        return response

    def test_classeur_des_etats(self):
        import openpyxl

        content = self._get(f"/finance/statements/export?kind=all&date_to={self.today.isoformat()}").content
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        self.assertEqual(book.sheetnames, ["Bilan", "Compte de résultat", "Tableau des flux", "Balance", "Flux mensuels"])
        header = [row for row in book["Bilan"].iter_rows(max_row=8, values_only=True)]
        self.assertTrue(any(row[0] == "Société" for row in header), "En-tête : société")
        self.assertTrue(any(row[0] == "Exporté le" for row in header), "En-tête : date d'export")

        status, body = self.call("GET", "/finance/statements/export?kind=all&format=csv", token=self.token)
        self.assertEqual(status, 400, "Un export CSV porte sur un seul état")

    def test_flux_mensuels_du_tableau_de_bord(self):
        """Flux mensuels exportés : le calcul du tableau de bord (monthly_cash), contrôle de chaque mois,
        écritures arrêtées à la date d'arrêté."""
        metrics = self.env["eguitra.finance.metrics"]
        fy_start = metrics.fiscal_year_bounds(self.company, self.today)[0]
        bank, receivable = self.bank.default_account_id, self.customer.property_account_receivable_id
        for day, amount in ((fy_start, 30_000_000), (fy_start + timedelta(days=1), 5_000_000)):
            self.env["account.move"].create({
                "move_type": "entry", "journal_id": self.bank.id, "date": day,
                "line_ids": [
                    Command.create({"name": "Règlement client", "account_id": bank.id, "debit": amount}),
                    Command.create({"name": "Règlement client", "account_id": receivable.id,
                                    "partner_id": self.customer.id, "credit": amount}),
                ],
            }).action_post()
        self.env.flush_all()  # visible des requêtes HTTP

        def rows(day):
            text = self._get(f"/finance/statements/export?kind=flux_mensuels&format=csv&date_to={day.isoformat()}")
            return list(csv.reader(io.StringIO(text.content.decode("utf-8-sig")), delimiter=";"))

        table = rows(self.today)
        header, body, totals = table[0], table[1:-1], table[-1]
        self.assertEqual(header[-1], "Contrôle")
        self.assertEqual(totals[0], "Période")
        months = [m for m in metrics.monthly_cash(self.company, self.today)["months"] if m["month"] <= f"{self.today:%Y-%m}"]
        self.assertEqual([r[0] for r in body], [m["month"] for m in months], "Mois de l'exercice jusqu'à l'arrêté")
        for row, month in zip(body, months, strict=True):
            self.assertEqual(row[header.index("Encaissements")], f"{month['receipts']:.0f}")
            self.assertEqual(row[header.index("Décaissements")], f"{month['payments']:.0f}")
            self.assertEqual(row[header.index("Trésorerie de fin de mois")], f"{month['end']:.0f}")
            self.assertEqual(row[header.index("CA HT du mois")], f"{month['revenue']:.0f}")
            self.assertEqual(row[header.index("Contrôle")], CASH_STATUS[month["status"]])

        first = rows(fy_start)
        self.assertEqual(len(first), 3, "Un mois, entre l'en-tête et les totaux")
        self.assertEqual(first[1][first[0].index("Encaissements")], "30000000",
                         "Le règlement du lendemain de l'arrêté n'est pas compté")

    def test_balance_classes_et_soldes(self):
        text = self._get("/finance/statements/export?kind=balance&format=csv").content.decode("utf-8-sig")
        self.assertIn("Solde débiteur", text)
        self.assertIn("Total classe 7", text)
        self.assertIn("Total classe 4", text)
        status, result = self.call("POST", "/finance/reports/run", {"report": "balance", "period": "year"}, token=self.token)
        self.assertEqual(status, 200, result)
        self.assertIn("classe", [c["key"] for c in result["columns"]])
        totals = result["totals"]
        self.assertAlmostEqual(totals["solde_debiteur"] - totals["solde_crediteur"], 0, 2,
                               msg="Soldes débiteurs = soldes créditeurs sur une balance équilibrée")
        self.assertGreater(totals["solde_debiteur"], 0)

    def test_grand_livre_et_compte_de_resultat(self):
        text = self._get("/finance/statements/export?kind=grand_livre&format=csv").content.decode("utf-8-sig")
        self.assertIn("Activité", text)
        self.assertIn("Report à nouveau", text)
        self.assertIn(f"Total {self.acc_services.code}", text)
        status, catalog = self.call("GET", "/finance/reports/catalog", token=self.token)
        ledger = next(r for g in catalog["groups"] for r in g["reports"] if r["key"] == "grand_livre")
        self.assertIn("activite", [c["key"] for c in ledger["columns"]])
        status, result = self.call("POST", "/finance/reports/run", {"report": "grand_livre", "period": "year"},
                                   token=self.token)
        self.assertEqual(status, 200, result)
        self.assertAlmostEqual(result["totals"]["debit"], result["totals"]["credit"], 2)

        status, income = self.call("GET", f"/finance/statements/income?date_to={self.today.isoformat()}", token=self.token)
        self.assertEqual(status, 200, income)
        revenue = next(line for line in income["lines"] if line["code"] == "XB")
        self.assertEqual(revenue["pct_revenue"], 100)
        self.assertEqual(revenue["month"], 200_000_000, "Colonne du mois")
        text = self._get("/finance/statements/export?kind=resultat&format=csv").content.decode("utf-8-sig")
        self.assertIn("% du CA", text)
