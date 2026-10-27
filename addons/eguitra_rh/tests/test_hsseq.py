import io
from datetime import date, timedelta

from dateutil.relativedelta import relativedelta

from odoo.exceptions import UserError
from odoo.tests import TransactionCase, tagged

from odoo.addons.eguitra_api.files import XLSX
from odoo.addons.eguitra_api.tests.common import ApiCase
from odoo.addons.eguitra_finance.tests.common import FinanceDataMixin

MANDATORY = ("PERMIS", "CONDUITE_DEFENSIVE", "MATIERES_DANGEREUSES")


@tagged("post_install", "-at_install")
class TestHsseq(FinanceDataMixin, TransactionCase):
    """Données RH lues par NIMOVA · HSSEQ : statuts seulement, jamais de détail médical."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.setup_finance_data()
        cls.day = date.today()
        cls.driver = cls.env["eguitra.transport.driver"].create({"name": "Mamadou Chauffeur", "company_id": cls.company.id})
        cls.emp = cls.env["eguitra.rh.employee"].create({
            "name": "Mamadou Chauffeur", "hire_date": date(2024, 1, 1), "base_salary": 3_000_000,
            "driver_id": cls.driver.id, "medical_fitness": "apte",
            "medical_fitness_until": cls.day + timedelta(days=200),
        })
        cls.Type = cls.env["eguitra.rh.qualification.type"]

    def _qual(self, code, expiry, emp=None):
        return self.env["eguitra.rh.qualification"].create({
            "employee_id": (emp or self.emp).id, "type_id": self.Type.search([("code", "=", code)]).id,
            "obtained_date": date(2023, 1, 1), "expiry_date": expiry})

    def test_referentiel(self):
        codes = set(self.Type.search([]).mapped("code"))
        self.assertTrue({"PERMIS", "CONDUITE_DEFENSIVE", "MATIERES_DANGEREUSES", "ACCES_SITE_MINIER", "SECOURISME"} <= codes)
        self.assertTrue(self.env.ref("eguitra_rh.leave_work_accident").work_accident)

    def test_verrou_v04(self):
        Employee = self.env["eguitra.rh.employee"]
        status = Employee._hsseq_driver_status(self.driver, self.day)
        self.assertTrue(status["blocking"], "Habilitations obligatoires absentes")
        self.assertEqual({q["code"] for q in status["qualifications"] if q["state"] == "absente"}, set(MANDATORY))
        for code in MANDATORY:
            self._qual(code, self.day + timedelta(days=365))
        status = Employee._hsseq_driver_status(self.driver, self.day)
        self.assertFalse(status["blocking"], status["reasons"])
        self.assertEqual(status["fitness"], "apte")
        self.assertNotIn("medical", " ".join(status), "Aucun champ médical exposé")

        self._qual("SECOURISME", self.day - timedelta(days=3))
        self.assertFalse(Employee._hsseq_driver_status(self.driver, self.day)["blocking"],
                         "Une habilitation non obligatoire expirée ne bloque pas le départ")
        expired = self.env["eguitra.rh.qualification"].search([("employee_id", "=", self.emp.id),
                                                               ("type_id.code", "=", "CONDUITE_DEFENSIVE")])
        expired.expiry_date = self.day - timedelta(days=1)
        status = Employee._hsseq_driver_status(self.driver, self.day)
        self.assertTrue(status["blocking"])
        self.assertIn("Conduite défensive : expirée.", status["reasons"])
        expired.expiry_date = self.day
        self.assertFalse(Employee._hsseq_driver_status(self.driver, self.day)["blocking"],
                         "Valide le jour même de son expiration ; V04 le lendemain")

        self.emp.medical_fitness = "inapte"
        status = Employee._hsseq_driver_status(self.driver, self.day)
        self.assertEqual((status["fitness"], status["blocking"]), ("inapte", True))
        self.emp.write({"medical_fitness": "apte", "medical_fitness_until": self.day - timedelta(days=1)})
        self.assertIn("Aptitude médicale échue.", Employee._hsseq_driver_status(self.driver, self.day)["reasons"])
        self.emp.write({"medical_fitness": "apte_reserve", "medical_fitness_until": self.day + timedelta(days=90)})
        status = Employee._hsseq_driver_status(self.driver, self.day)
        self.assertEqual((status["fitness"], status["blocking"]), ("apte_reserve", False))
        self.assertIn("Aptitude médicale avec réserve.", status["warnings"])
        self.emp.medical_fitness = False
        status = Employee._hsseq_driver_status(self.driver, self.day)
        self.assertFalse(status["blocking"], "Aptitude non renseignée : avertissement, pas motif V04")
        self.assertIn("Aptitude médicale non renseignée.", status["warnings"])

        other = self.env["eguitra.transport.driver"].create({"name": "Sans dossier", "company_id": self.company.id})
        self.assertTrue(Employee._hsseq_driver_status(other, self.day)["blocking"], "Chauffeur sans dossier salarié")

    def test_alertes_j30_j7(self):
        Employee = self.env["eguitra.rh.employee"]
        self._qual("ACCES_SITE_MINIER", self.day + timedelta(days=20))
        soon = self._qual("SECOURISME", self.day + timedelta(days=5))
        self._qual("PERMIS", self.day + timedelta(days=90))
        rows = {r["code"]: r for r in Employee._hsseq_expiring(self.company, 30, self.day)}
        self.assertEqual((rows["ACCES_SITE_MINIER"]["days_left"], rows["SECOURISME"]["days_left"]), (20, 5))
        self.assertNotIn("PERMIS", rows, "Hors de l'horizon de 30 jours")
        self._qual("SECOURISME", self.day + timedelta(days=700))
        rows = {r["code"]: r for r in Employee._hsseq_expiring(self.company, 30, self.day)}
        self.assertNotIn("SECOURISME", rows, "Renouvelée : l'ancienne habilitation ne déclenche plus d'alerte")
        self.assertEqual(soon._state(self.day), "a_renouveler")
        with self.assertRaises(UserError):
            self.env["eguitra.rh.qualification"].create({
                "employee_id": self.emp.id, "type_id": self.Type.search([("code", "=", "PERMIS")]).id,
                "obtained_date": self.day, "expiry_date": self.day - timedelta(days=1)})

    def test_heures_et_accidents_du_travail(self):
        Employee = self.env["eguitra.rh.employee"]
        sept = date(2026, 9, 1)
        payroll = self.env["eguitra.rh.payroll"].prepare(self.company, sept)
        slip = payroll.payslip_ids.filtered(lambda s: s.employee_id == self.emp)
        self.assertEqual(slip.worked_days, slip.month_days)
        # Quatre yeux : la paie préparée par un compte est validée par la direction.
        valideur = self.env["res.users"].with_context(no_reset_password=True).create(
            {"name": "Direction Test", "login": "direction.hsseq"})
        valideur._eguitra_set_profiles(["direction"])
        payroll.with_user(valideur).action_validate()
        settings = self.env["eguitra.rh.settings"]._get(self.company)
        worked = Employee._hsseq_worked_hours(self.company, sept, date(2026, 9, 30))
        hours = worked["hours"]
        self.assertAlmostEqual(hours, round(settings.monthly_hours * len(payroll.payslip_ids), 1), delta=0.2)
        self.assertEqual((worked["basis"], [m["basis"] for m in worked["months"]]), ("paid", ["paid"]),
                         "Heures payées d'une paie validée")
        october = Employee._hsseq_worked_hours(self.company, date(2026, 10, 1), date(2026, 10, 31))
        self.assertEqual(october["basis"], "presence", "Sans paie validée : estimation par la présence")

        accident = self.env["eguitra.rh.leave"].create({
            "employee_id": self.emp.id, "type_id": self.env.ref("eguitra_rh.leave_work_accident").id,
            "date_from": date(2026, 9, 10), "date_to": date(2026, 9, 12)})
        accident.action_approve(force=True)
        accidents = Employee._hsseq_work_accidents(self.company, sept, date(2026, 9, 30))
        self.assertEqual([a["lost_days"] for a in accidents], [3], "Jours perdus calendaires")
        rates = Employee._hsseq_indicators(self.company, sept, date(2026, 9, 30))
        self.assertEqual((rates["accidents"], rates["lost_days"]), (1, 3))
        self.assertAlmostEqual(rates["frequency_rate"], 1_000_000 / hours, delta=0.1)
        self.assertAlmostEqual(rates["severity_rate"], 3 * 1000 / hours, delta=0.1)


@tagged("post_install", "-at_install")
class TestHsseqApi(ApiCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.rh_user = cls._create_user("rh.securite", "rh")
        cls.worker_user = cls._create_user("chauffeur.mobile", "exploitation")
        company = cls.rh_user.company_id
        cls.driver = cls.env["eguitra.transport.driver"].create({"name": "Chauffeur Mobile", "company_id": company.id})
        cls.worker = cls.env["eguitra.rh.employee"].create({
            "name": "Chauffeur Mobile", "hire_date": date(2024, 1, 1), "driver_id": cls.driver.id,
            "user_id": cls.worker_user.id, "company_id": company.id, "medical_fitness": "apte"})
        cls.own = cls.env["eguitra.rh.employee"].create({
            "name": "RH Sécurité", "hire_date": date(2024, 1, 1), "user_id": cls.rh_user.id, "company_id": company.id})

    def test_habilitations_par_la_rh(self):
        token = self.open_session("rh.securite", app="rh")["access_token"]
        status, meta = self.call("GET", "/rh/meta", token=token)
        permis = next(t for t in meta["qualification_types"] if t["code"] == "CONDUITE_DEFENSIVE")
        status, qual = self.call("POST", "/rh/qualifications", {
            "employee_id": self.worker.id, "type_id": permis["id"], "obtained_date": "2026-01-15"}, token=token)
        self.assertEqual(status, 200, qual)
        self.assertEqual(qual["expiry_date"], "2028-01-14", "Expiration proposée par la durée de validité")
        status, matrix = self.call("GET", "/rh/qualifications/matrix", token=token)
        self.assertEqual(status, 200, matrix)
        row = next(r for r in matrix["rows"] if r["employee"]["id"] == self.worker.id)
        self.assertEqual(row["cells"]["CONDUITE_DEFENSIVE"]["state"], "valide")
        self.assertTrue(row["blocking"], "Permis et matières dangereuses manquent")
        status, _b = self.call("POST", "/rh/qualifications", {"employee_id": self.own.id, "type_id": permis["id"]},
                               token=token)
        self.assertEqual(status, 403, "Personne ne gère ses propres habilitations")
        status, ppe = self.call("POST", "/rh/ppe", {"employee_id": self.worker.id, "item": "casque"}, token=token)
        self.assertEqual((status, ppe["item"]), (200, "casque"), ppe)
        status, safety = self.call("GET", "/rh/safety", token=token)
        self.assertEqual(status, 200, safety)
        self.assertEqual(safety["drivers"], 1)

    def test_conducteur_consulte_ses_habilitations(self):
        token = self.open_session("chauffeur.mobile", app="rh")["access_token"]
        status, me = self.call("GET", "/rh/me", token=token)
        self.assertEqual(status, 200, me)
        self.assertEqual(me["fitness"]["state"], "apte")
        self.assertEqual({q["code"] for q in me["qualifications"] if q["state"] == "absente"}, set(MANDATORY))
        status, _b = self.call("GET", "/rh/qualifications/matrix", token=token)
        self.assertEqual(status, 403, "La matrice est réservée à la gestion")

    def _classeur(self, habilitations, aptitudes):
        import openpyxl

        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Habilitations"
        sheet.append(["Matricule", "Salarié", "Société", "Code habilitation", "Habilitation", "Obtenue le",
                      "Expire le", "Référence", "Délivrée par"])
        for row in habilitations:
            sheet.append(row)
        sheet = book.create_sheet("Aptitude")
        sheet.append(["Matricule", "Salarié", "Société", "Aptitude", "Visite le", "Valable jusqu'au"])
        for row in aptitudes:
            sheet.append(row)
        out = io.BytesIO()
        book.save(out)
        return out.getvalue()

    def _importer(self, token, content):
        response = self.url_open("/api/v1/rh/qualifications/import", method="POST", timeout=60,
                                 files={"file": ("reprise.xlsx", io.BytesIO(content), XLSX)},
                                 headers={"Authorization": f"Bearer {token}", "X-Eguitra-App": "test"})
        return response.status_code, response.json()

    def test_reprise_par_classeur(self):
        import openpyxl

        token = self.open_session("rh.securite", app="rh")["access_token"]
        response = self.url_open("/api/v1/rh/qualifications/template", timeout=60,
                                 headers={"Authorization": f"Bearer {token}", "X-Eguitra-App": "test"})
        self.assertEqual(response.status_code, 200, response.text[:300])
        book = openpyxl.load_workbook(io.BytesIO(response.content), read_only=True)
        self.assertEqual(book.sheetnames, ["Habilitations", "Aptitude", "Notice"])
        rows = list(book["Habilitations"].iter_rows(values_only=True))
        codes = {r[3] for r in rows[1:] if r[0] == self.worker.matricule}
        self.assertTrue(set(MANDATORY) <= codes, "Une rangée par type d'habilitation pour chaque conducteur")
        self.assertNotIn(self.own.matricule, {r[0] for r in rows[1:]}, "Conducteurs seulement par défaut")

        m, day = self.worker.matricule, date.today()
        obtained = day - timedelta(days=60)
        content = self._classeur([
            [m, None, None, "PERMIS", None, date(2020, 3, 1), date(2030, 3, 1), "P-1", "Préfecture"],
            [m.lower(), None, None, "conduite_defensive", None, obtained.strftime("%d/%m/%Y"), None, None, None],
            [m, None, None, "MATIERES_DANGEREUSES", None, "2025-06-01", "2030-05-31", None, None],
            [m, None, None, "SECOURISME", None, None, None, None, None],
        ], [[m, None, None, "Apte avec réserve", day - timedelta(days=30), day + timedelta(days=330)]])
        status, result = self._importer(token, content)
        self.assertEqual(status, 200, result)
        self.assertEqual(result, {"created": 3, "updated": 0, "fitness": 1, "unchanged": 0})
        self.env.invalidate_all()
        defensive = self.worker.qualification_ids.filtered(lambda q: q.type_id.code == "CONDUITE_DEFENSIVE")
        self.assertEqual(defensive.expiry_date, obtained + relativedelta(months=24, days=-1),
                         "Expiration calculée par la durée du type")
        self.assertEqual(self.worker.medical_fitness, "apte_reserve")
        status, matrix = self.call("GET", "/rh/qualifications/matrix", token=token)
        row = next(r for r in matrix["rows"] if r["employee"]["id"] == self.worker.id)
        self.assertFalse(row["blocking"], row["reasons"])

        status, again = self._importer(token, content)
        self.assertEqual((status, again), (200, {"created": 0, "updated": 0, "fitness": 0, "unchanged": 4}),
                         "Réimporter le même classeur ne crée rien")

        before = self.env["eguitra.rh.qualification"].search_count([])
        status, body = self._importer(token, self._classeur([
            ["ZZ-9999", None, None, "PERMIS", None, "01/01/2025", None, None, None],
            [m, None, None, "INCONNU", None, "01/01/2025", None, None, None],
            [m, None, None, "SECOURISME", None, "31/02/2025", None, None, None],
            [self.own.matricule, None, None, "SECOURISME", None, "01/01/2025", None, None, None],
            [m, None, None, "SECOURISME", None, "01/01/2025", None, None, None],
        ], [[m, None, None, "Bon pour le service", None, None]]))
        self.assertEqual(status, 422, body)
        self.assertEqual(len(body["errors"]), 5, body["errors"])
        self.assertIn("rien n'est enregistré", body["message"])
        self.assertEqual(self.env["eguitra.rh.qualification"].search_count([]), before, "Tout ou rien")
