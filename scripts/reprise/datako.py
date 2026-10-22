"""Conversion des exports de Datakö Fleet (logiciel de transport utilisé avant NIMOVA) en fichiers de
reprise du transport, au format lu par reprise.py (modèles dans modeles/).

    python -I scripts/reprise/datako.py livraisons.xlsx [autre_periode.xlsx …] dossier_de_sortie
        [--tarifs=tarifs.csv] [--routes=routes_datako.csv] [--carburant=carburant.csv]
        [--depuis=AAAA-MM-JJ] [--facturees-jusqu-au=AAAA-MM-JJ]

Entrée principale : le classeur de l'écran « Livraisons » de Datakö (bouton « Exporter », feuille
« Livraisons », 21 colonnes), un ou plusieurs fichiers (une période chacun) ; une rotation présente
dans deux fichiers n'est reprise qu'une fois (le dernier fichier l'emporte). Cet export arrondit le
tarif au franc et ne porte ni conducteur, ni coûts, ni encaissements ; trois fichiers facultatifs,
recopiés des écrans de configuration de Datakö (séparateur point-virgule, décimales à la virgule ou
au point), le complètent :

- --tarifs : `depot;destination;ville;tarif_l` (écran « Tarifs de transport », tarif appliqué en
  GNF par litre, deux décimales). Il rend au tarif ses décimales et retrouve la destination de la
  route, que l'export ne donne pas (il ne donne que la ville) : parmi les destinations du dépôt et de
  la ville, celle dont le tarif arrondi est celui de l'export.
- --routes : `depot;destination;distance_km;conso_l;prime;peage;frais_chargement` (écran « Routes »).
  Il donne la distance des routes et les coûts déclarés des rotations, tels que Datakö les compte :
  prime du chauffeur, péages, frais de chargement et carburant (consommation de la route × prix du
  mois de chargement).
- --carburant : `mois;prix_l` (AAAA-MM, prix du gasoil appliqué par Datakö ce mois-là). Le carburant
  d'une rotation est figé au prix de son mois ; sans prix du mois, il n'est pas repris (signalé).

Rotations reprises : les livraisons « Livré ». Celles « En cours » à la date de l'export se saisissent
dans NIMOVA après la bascule ; les « Annulé » ne se reprennent pas (les deux sont listées). Une
rotation livrée est marquée facturée dans l'ancien système (sa facture est dans la comptabilité
reprise), sauf après la date --facturees-jusqu-au : elle sera facturée par NIMOVA. --depuis écarte
les rotations chargées avant une date (période déjà reprise d'une autre source, le classeur 2026).

Volumes en m³ (litres ÷ 1 000), tarif par m³ (tarif par litre × 1 000). Le numéro de rotation
Datakö, le n° de tournée du client et le site de livraison passent dans les observations.

Les exports sont des données du client, lues sans rien exécuter ; les fichiers produits ne se
versionnent pas. Fichiers produits : routes, ensembles, rotations (CSV) et rapport_conversion.txt
(totaux par mois, écarts de chiffre d'affaires dus à l'arrondi du tarif, clients à déclarer dans
tiers.csv, points à vérifier).
"""

import csv
import math
import re
import sys
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

HEADERS = {
    "routes": "nom origine destination distance_km distance_ar_km duree_jours tarif_m3",
    "ensembles": "code immatriculation marque capacite_litres chauffeur annee mise_en_service etat immobilisation",
    "rotations": "date_chargement date_livraison ensemble chauffeur client route produit bl volume_charge volume_livre "
                 "tarif carburant peages frais_chauffeur entretien taxes autres facturee observations date_retour "
                 "litres km pneus",
}
# Colonnes de l'export « Livraisons » de Datakö (libellés sans accents ni casse).
COLUMNS = {
    "dossier": "dossier", "date chargement": "loaded_on", "client": "client", "camion": "truck",
    "n rotation": "number", "n bl": "bl", "n tournee client": "tour", "produit": "product",
    "depot de chargement": "depot", "ville": "city", "site de livraison": "site",
    "qte chargee (l)": "loaded", "qte livree (l)": "delivered", "qte manquante (l)": "shortage",
    "ca (gnf)": "revenue", "tarif/l (gnf)": "rate", "statut livraison": "status",
    "confirmation destinataire": "confirmation", "date livraison": "delivered_on", "encaissement": "payment",
    "motif annulation": "cancel_reason",
}
REQUIRED = ("loaded_on", "client", "truck", "number", "depot", "city", "loaded", "delivered", "rate", "status")
MONTHS = {"janv": 1, "fevr": 2, "mars": 3, "avr": 4, "mai": 5, "juin": 6, "juil": 7, "aout": 8, "sept": 9,
          "oct": 10, "nov": 11, "dec": 12}
# Produits de Datakö -> codes des produits de la plateforme.
PRODUCTS = (("gasoil", "gasoil"), ("gazole", "gasoil"), ("essence", "essence"), ("mazout", "fuel"),
            ("hfo", "fuel"), ("fuel", "fuel"), ("jet", "jet"), ("petrole", "petrole"))


# ─── Outils ───────────────────────────────────────────────────────────────

def plain(text):
    """Minuscules sans accents ni espaces superflus, pour comparer deux libellés."""
    text = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", text).strip().lower()


def number(value):
    """Nombre d'une cellule ou d'un texte (« 40 000 », « 1 234,56 ») ; None si vide."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = re.sub(r"[^\d,.\-]", "", str(value))
    if text in ("", "-"):
        return None
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def day(value):
    """Date d'une cellule : date du tableur, « 10 oct. 2026 », JJ/MM/AAAA ou AAAA-MM-JJ ; None sinon."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = plain(value)
    m = re.fullmatch(r"(\d{1,2}) ([a-z]+)\.? (\d{4})", text)
    month = m and (MONTHS.get(m.group(2)[:4]) or MONTHS.get(m.group(2)[:3]))
    if month:
        return date(int(m.group(3)), month, int(m.group(1)))
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", text)
    if m:
        return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", text[:10])
    return date(int(m.group(1)), int(m.group(2)), int(m.group(3))) if m else None


def fr(d):
    return d.strftime("%d/%m/%Y") if d else ""


def fmt(value, digits=0):
    if value is None:
        return ""
    if digits:
        return f"{value:.{digits}f}".rstrip("0").rstrip(".")
    return str(int(round(value)))


def money(value):
    return f"{value:,.0f}".replace(",", " ")


def read_csv(path, columns):
    """Lignes d'un fichier complémentaire (point-virgule), colonnes attendues vérifiées."""
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f, delimiter=";"))
    missing = [c for c in columns if rows and c not in rows[0]]
    if missing:
        sys.exit(f"{path} : colonne(s) manquante(s) : {', '.join(missing)}.")
    return rows


def read_export(path):
    """Livraisons d'un export Datakö : une liste de dictionnaires aux clés de COLUMNS."""
    try:
        import openpyxl
    except ImportError:
        sys.exit("Le module openpyxl est nécessaire (pip install openpyxl).")
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = book["Livraisons"] if "Livraisons" in book.sheetnames else book.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = [COLUMNS.get(plain(h)) for h in next(rows, ())]
    missing = [k for k in REQUIRED if k not in header]
    if missing:
        sys.exit(f"{path} : ce n'est pas un export « Livraisons » de Datakö (colonnes absentes : {', '.join(missing)}).")
    result = []
    for values in rows:
        if not any(v not in (None, "") for v in values):
            continue
        result.append({key: value for key, value in zip(header, values) if key})
    book.close()
    return result


# ─── Conversion ───────────────────────────────────────────────────────────

class Export:
    def __init__(self, tariffs=None, routes=None, fuel=None, since=None, billed_until=None):
        self.tariffs = defaultdict(list)  # (dépôt, ville) -> [(destination, tarif_l)]
        for r in tariffs or ():
            rate = number(r["tarif_l"])
            if rate is not None:
                self.tariffs[(plain(r["depot"]), plain(r["ville"]))].append((r["destination"].strip(), rate))
        self.route_costs = {(plain(r["depot"]), plain(r["destination"])): r for r in routes or ()}
        self.fuel = {r["mois"].strip(): number(r["prix_l"]) for r in fuel or ()}
        self.since, self.billed_until = since, billed_until
        self.rows = {name: [] for name in HEADERS}
        self.report, self.warnings = [], []

    def route_of(self, depot, city, rate):
        """(destination, tarif par litre, exact) d'une livraison : la destination du dépôt et de la ville
        dont le tarif arrondi est celui de l'export ; à défaut, la ville et le tarif arrondi."""
        candidates = self.tariffs.get((plain(depot), plain(city)), [])
        if rate is not None:
            matching = [c for c in candidates if abs(c[1] - rate) < 0.5 + 1e-9]
            if len(matching) > 1:
                # Même tarif arrondi pour deux destinations : la destination qui porte le nom de la ville.
                matching = [c for c in matching if plain(c[0]) == plain(city)] or matching
            if len(matching) == 1:
                return matching[0][0], matching[0][1], True
        if len(candidates) == 1 and rate is None:
            return candidates[0][0], candidates[0][1], True
        return city, rate, False

    def convert(self, deliveries):
        by_number = {}
        for row in deliveries:
            key = str(row.get("number") or "").strip()
            if key in by_number:
                self.warnings.append(f"Rotation {key} présente dans plusieurs exports : la dernière lue est reprise.")
            by_number[key or f"sans-numero-{len(by_number)}"] = row
        routes, trucks, clients = {}, set(), defaultdict(int)
        months = defaultdict(lambda: {"count": 0, "m3": 0.0, "export": 0.0, "recalc": 0.0})
        pending, cancelled, before, rounded, no_price = [], [], 0, set(), set()
        for key, r in sorted(by_number.items(), key=lambda kv: (day(kv[1].get("loaded_on")) or date.min, kv[0])):
            status = plain(r.get("status"))
            loaded_on, delivered_on = day(r.get("loaded_on")), day(r.get("delivered_on"))
            if status.startswith("annul"):
                reason = str(r.get("cancel_reason") or "").strip()
                cancelled.append(f"{key}{' (' + reason + ')' if reason else ''}")
                continue
            if not status.startswith("livr"):
                pending.append(key)
                continue
            if not loaded_on:
                self.warnings.append(f"Rotation {key} : date de chargement illisible « {r.get('loaded_on')} », non reprise.")
                continue
            if self.since and loaded_on < self.since:
                before += 1
                continue
            depot, city = str(r.get("depot") or "").strip(), str(r.get("city") or "").strip()
            destination, rate_l, exact = self.route_of(depot, city, number(r.get("rate")))
            if not exact:
                rounded.add(f"{depot} - {city}")
            if rate_l is None:
                self.warnings.append(f"Rotation {key} : tarif absent, non reprise.")
                continue
            name = f"{depot} - {destination}"
            costs = self.route_costs.get((plain(depot), plain(destination)))
            routes.setdefault(name, (depot, destination, rate_l, costs))
            truck, client = str(r.get("truck") or "").strip(), str(r.get("client") or "").strip()
            trucks.add(truck)
            clients[client] += 1
            if not client:
                self.warnings.append(f"Rotation {key} : sans client dans Datakö ; à compléter avant la reprise.")
            loaded, delivered = number(r.get("loaded")) or 0, number(r.get("delivered")) or 0
            fuel_l = fuel = peages = prime = other = None
            if costs:
                fuel_l = number(costs.get("conso_l"))
                price = self.fuel.get(f"{loaded_on:%Y-%m}")
                if fuel_l and price:
                    fuel = fuel_l * price
                elif fuel_l:
                    no_price.add(f"{loaded_on:%Y-%m}")
                peages, prime, other = (number(costs.get(c)) for c in ("peage", "prime", "frais_chargement"))
            billed = not self.billed_until or (delivered_on or loaded_on) <= self.billed_until
            notes = [f"Datakö {key}"]
            if r.get("tour"):
                notes.append(f"tournée {r['tour']}")
            if r.get("site"):
                notes.append(f"site {str(r['site']).strip()}")
            if plain(r.get("dossier")).startswith("incomplet"):
                notes.append("dossier incomplet dans Datakö")
            product = next((code for word, code in PRODUCTS if word in plain(r.get("product"))), "gasoil")
            self.rows["rotations"].append([
                fr(loaded_on), fr(delivered_on), truck, "", client, name, product, str(r.get("bl") or "").strip(),
                fmt(loaded / 1000, 3), fmt(delivered / 1000, 3), fmt(rate_l * 1000, 2), fmt(fuel), fmt(peages),
                fmt(prime), "", "0", fmt(other or None), "oui" if billed else "", " ; ".join(notes), "",
                fmt(fuel_l, 1), "", "",
            ])
            m = months[f"{(delivered_on or loaded_on):%Y-%m}"]
            m["count"] += 1
            m["m3"] += delivered / 1000
            m["export"] += number(r.get("revenue")) or 0
            m["recalc"] += math.floor(delivered * rate_l + 0.5)  # arrondi de Datakö et du noyau : demi vers le haut
        for name, (depot, destination, rate_l, costs) in sorted(routes.items()):
            if self.route_costs and not costs:
                self.warnings.append(f"Route « {name} » absente de --routes : ni distance ni coûts déclarés.")
            distance = number(costs.get("distance_km")) if costs else None
            self.rows["routes"].append([name, depot, destination, fmt(distance), "", "", fmt(rate_l * 1000, 2)])
        for truck in sorted(trucks):
            self.rows["ensembles"].append([truck, "", "", "", "", "", "", "service", ""])
        # Compte rendu.
        total = sum(m["count"] for m in months.values())
        self.report.append(f"Rotations livrées reprises : {total} ; routes : {len(routes)} ; ensembles : {len(trucks)}.")
        if before:
            self.report.append(f"Rotations chargées avant le {fr(self.since)} écartées (--depuis) : {before}.")
        self.report.append("Par mois de livraison : rotations, volume livré, chiffre d'affaires de l'export et recalculé "
                           "(volume livré × tarif repris) :")
        for month, m in sorted(months.items()):
            gap = m["recalc"] - m["export"]
            self.report.append(f"  {month} : {m['count']} rotation(s), {fmt(m['m3'], 3)} m³, {money(m['export'])} GNF "
                               f"dans l'export, {money(m['recalc'])} GNF recalculés (écart {money(gap)}).")
        self.report.append("Clients à déclarer dans tiers.csv sous ce nom exact (ou à renommer ici) : "
                           + ", ".join(f"« {c} » ({n})" for c, n in sorted(clients.items()) if c) + ".")
        if pending:
            self.report.append(f"En cours dans Datakö, à saisir dans NIMOVA après la bascule : {', '.join(pending)}.")
        if cancelled:
            self.report.append(f"Annulées, non reprises : {', '.join(cancelled)}.")
        if rounded:
            self.warnings.append("Tarif arrondi au franc de l'export (destination non retrouvée dans --tarifs) pour : "
                                 + ", ".join(sorted(rounded)) + ".")
        if no_price:
            self.warnings.append(f"Prix du carburant absent (--carburant) pour {', '.join(sorted(no_price))} : carburant "
                                 "des rotations chargées ces mois-là non repris.")
        if not self.route_costs:
            self.warnings.append("Sans --routes : aucun coût déclaré ni distance ; la rentabilité ne portera que les "
                                 "coûts comptabilisés.")
        self.warnings.append("Conducteurs absents de l'export Datakö : rotations sans chauffeur (à compléter s'il le faut).")
        self.warnings.append("Ensembles : code = immatriculation de Datakö ; compléter tracteur, citerne, marque et "
                             "capacité dans ensembles.csv (ou dans l'écran Flotte).")
        if not self.billed_until:
            self.warnings.append("Toutes les rotations livrées sont marquées facturées dans l'ancien système "
                                 "(--facturees-jusqu-au pour laisser les plus récentes à facturer par NIMOVA).")

    def write(self, folder):
        folder.mkdir(parents=True, exist_ok=True)
        for name, header in HEADERS.items():
            with open(folder / f"{name}.csv", "w", encoding="utf-8-sig", newline="") as f:
                writer = csv.writer(f, delimiter=";", lineterminator="\n")
                writer.writerow(header.split())
                writer.writerows(self.rows[name])
        text = ["Conversion des exports Datakö : compte rendu", ""] + [f"- {line}" for line in self.report]
        text += ["", f"Points à vérifier ({len(self.warnings)}) :"] + [f"- {w}" for w in self.warnings or ["aucun."]]
        (folder / "rapport_conversion.txt").write_text("\n".join(text) + "\n", encoding="utf-8")
        return text


def main(argv):
    args = [a for a in argv if not a.startswith("--")]
    options = dict(a[2:].split("=", 1) if "=" in a else (a[2:], "1") for a in argv if a.startswith("--"))
    if len(args) < 2:
        sys.exit("Usage : python -I scripts/reprise/datako.py livraisons.xlsx [autre.xlsx …] dossier_de_sortie "
                 "[--tarifs=tarifs.csv] [--routes=routes_datako.csv] [--carburant=carburant.csv] "
                 "[--depuis=AAAA-MM-JJ] [--facturees-jusqu-au=AAAA-MM-JJ]")
    *sources, target = args
    dates = {}
    for opt in ("depuis", "facturees-jusqu-au"):
        if opt in options:
            dates[opt] = day(options[opt])
            if not dates[opt]:
                sys.exit(f"--{opt} : date attendue au format AAAA-MM-JJ.")
    export = Export(
        tariffs=read_csv(options["tarifs"], ("depot", "destination", "ville", "tarif_l")) if "tarifs" in options else None,
        routes=read_csv(options["routes"], ("depot", "destination", "distance_km", "conso_l", "prime", "peage",
                                            "frais_chargement")) if "routes" in options else None,
        fuel=read_csv(options["carburant"], ("mois", "prix_l")) if "carburant" in options else None,
        since=dates.get("depuis"), billed_until=dates.get("facturees-jusqu-au"),
    )
    deliveries = []
    for source in sources:
        deliveries += read_export(Path(source))
    export.convert(deliveries)
    print("\n".join(export.write(Path(target))))
    print(f"\nFichiers écrits dans {target}.")


if __name__ == "__main__":
    main(sys.argv[1:])
